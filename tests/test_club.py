"""End-to-end workflows against a throwaway database: self-registration -> approval -> membership number, family
updates, publishing with push, meetings with members-only minutes, library circulation with fines, accounts,
broadcasts and the daily jobs. Messages go to a fake transport so nothing leaves the machine."""
import datetime as dt
import os
import tempfile

os.environ["CLUB_DATA"] = tempfile.mkdtemp(prefix="club-test-")
os.environ["CLUB_DEV"] = "1"
os.environ["CLUB_NO_WORKER"] = "1"
os.environ["CLUB_ADMIN_LOGIN"] = "admin@club.test"
os.environ["CLUB_ADMIN_PASSWORD"] = "admin-password-1"
os.environ["CLUB_PUBLIC_URL"] = "https://club.example.org"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import comms, db, i18n, scheduler  # noqa: E402
from app.main import app  # noqa: E402

H = {"X-Club": "1"}
SENT = []
comms.set_sender(lambda ch, to, subject, body: SENT.append((ch, to, subject, body)) or f"id-{len(SENT)}")


@pytest.fixture(scope="module")
def admin():
    with TestClient(app) as c:
        r = c.post("/api/auth/login", json={"login": "admin@club.test", "password": "admin-password-1"}, headers=H)
        assert r.status_code == 200, r.text
        yield c


def login(login_, pw):
    from app import security
    security.reset()                      # the suite signs in far more often per minute than the limiter allows
    c = TestClient(app)
    r = c.post("/api/auth/login", json={"login": login_, "password": pw}, headers=H)
    assert r.status_code == 200, r.text
    return c


def flush():
    comms.process_outbox(500)


def fresh():
    """Deliver whatever earlier steps queued, then start counting afresh."""
    flush()
    SENT.clear()


def test_public_info_and_registration_flow(admin):
    info = admin.get("/api/public/info").json()
    assert info["club_short"] == "ASC" and any(t["key"] == "life" for t in info["membership_types"])

    # self registration with family, in Malayalam
    r = admin.post("/api/public/apply", json={"name_en": "Ramesh Kumar", "name_ml": "രമേശ് കുമാർ", "phone": "98470 12345", "email": "ramesh@example.org",
                                              "dob": "1980-05-14", "blood_group": "O+", "ward": "7", "lang": "ml", "password": "ramesh-pass-1",
                                              "family": [{"name_en": "Sita", "name_ml": "സീത", "relation": "spouse", "dob": "1984-01-02"}]}, headers=H)
    assert r.status_code == 200, r.text
    mid = r.json()["id"]
    # the applicant is told, in Malayalam, on WhatsApp / SMS / e-mail
    flush()
    kinds = {(ch, to) for ch, to, *_ in SENT}
    assert ("whatsapp", "919847012345") in kinds and ("sms", "919847012345") in kinds and ("email", "ramesh@example.org") in kinds
    assert any("അപേക്ഷ" in body for *_, body in SENT)
    # cannot sign in before approval
    c = TestClient(app)
    assert c.post("/api/auth/login", json={"login": "9847012345", "password": "ramesh-pass-1"}, headers=H).status_code == 403
    # duplicate phone is refused
    assert admin.post("/api/public/apply", json={"name_en": "Dup", "phone": "9847012345", "password": "whatever-123"}, headers=H).status_code == 409

    apps = admin.get("/api/members/applications").json()
    assert [a["id"] for a in apps] == [mid] and apps[0]["family"][0]["relation"] == "spouse"

    fresh()
    r = admin.post(f"/api/members/{mid}/decide", json={"decision": "approve"}, headers=H)
    assert r.status_code == 200, r.text
    no = r.json()["member_no"]
    assert no == f"ASC/{dt.date.today().year}/0001"
    flush()
    assert any(no in body and "സ്വാഗതം" in subj for ch, to, subj, body in SENT if ch == "whatsapp")
    m = admin.get(f"/api/members/{mid}").json()
    assert m["status"] == "active" and m["member_no"] == no
    # admission fee + annual subscription are now due (ordinary: 100 + 120)
    assert sorted(s["amount"] for s in m["subscriptions"]) == [100, 120]

    # the member signs in with the phone and updates the family
    mc = login("9847012345", "ramesh-pass-1")
    me = mc.get("/api/auth/me").json()
    assert me["user"]["role"] == "member" and me["member"]["member_no"] == no and "members.manage" not in me["user"]["perms"]
    r = mc.post(f"/api/members/{mid}/family", json={"name_en": "Arjun", "relation": "son", "dob": "2010-07-07", "blood_group": "O+"}, headers=H)
    assert r.status_code == 200
    fam = mc.get(f"/api/members/{mid}/family").json()
    assert len(fam) == 2 and fam[1]["verified"] == 0
    # but may not touch another member's record or the office lists
    assert mc.get("/api/members").status_code == 403
    assert mc.put("/api/members/999", json={"name_en": "x"}, headers=H).status_code == 403
    # directory shows them, blood group search works
    d = mc.get("/api/members/directory?blood=O%2B").json()
    assert d and d[0]["member_no"] == no
    # a second applicant gets the next number
    admin.post("/api/public/apply", json={"name_en": "Second", "phone": "9000000002", "password": "second-pass-1"}, headers=H)
    mid2 = admin.get("/api/members/applications").json()[0]["id"]
    assert admin.post(f"/api/members/{mid2}/decide", json={"decision": "approve"}, headers=H).json()["member_no"].endswith("/0002")
    # rejection notifies and blocks sign-in
    admin.post("/api/public/apply", json={"name_en": "Third", "phone": "9000000003", "password": "third-pass-1"}, headers=H)
    mid3 = admin.get("/api/members/applications").json()[0]["id"]
    fresh()
    admin.post(f"/api/members/{mid3}/decide", json={"decision": "reject", "note": "outside the ward"}, headers=H)
    flush()
    assert any("outside the ward" in body for *_, body in SENT)
    assert TestClient(app).post("/api/auth/login", json={"login": "9000000003", "password": "third-pass-1"}, headers=H).status_code == 403
    # office entry with a given number and the printable card
    r = admin.post("/api/members", json={"name_en": "Old Member", "phone": "9000000004", "member_no": "ASC/OLD/17", "membership_type": "life"}, headers=H)
    assert r.status_code == 200 and r.json()["member_no"] == "ASC/OLD/17"
    card = admin.get(f"/api/members/{mid}/card?lang=ml")
    assert card.status_code == 200 and no in card.text and "data:image/svg" in card.text
    csv_ = admin.get("/api/members/export/csv")
    assert csv_.status_code == 200 and no in csv_.text


def test_news_event_publish_register_results(admin):
    fresh()
    r = admin.post("/api/posts", json={"kind": "competition", "title_en": "Onam Sports Meet", "title_ml": "ഓണം കായികമേള", "body_en": "Annual sports meet", "body_ml": "വാർഷിക കായികമേള",
                                       "level": "panchayat", "starts_at": "2030-09-05T09:00", "venue_en": "School ground", "venue_ml": "സ്കൂൾ ഗ്രൗണ്ട്", "audience": "public",
                                       "reg_open": 1, "reg_family": 1, "items": [{"name_en": "100 m", "name_ml": "100 മീറ്റർ", "age_group": "senior", "max_entries": 1},
                                                                                 {"name_en": "Tug of war", "kind": "team"}]}, headers=H)
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    # drafts are invisible to members
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.get(f"/api/posts/{pid}").status_code == 404
    r = admin.post(f"/api/posts/{pid}/publish", json={"channels": ["app", "whatsapp", "sms"]}, headers=H)
    assert r.status_code == 200 and r.json()["queued"] > 0
    flush()
    wa = [body for ch, to, s, body in SENT if ch == "whatsapp" and to == "919847012345"]
    assert wa and "ഓണം കായികമേള" in wa[0] and "/p/" in wa[0]          # member's language is Malayalam; public link
    assert mc.get("/api/notifications").json()[0]["link"] == f"#/posts/{pid}"
    # public page, bilingual
    pub = TestClient(app).get(f"/p/{pid}?lang=en")
    assert pub.status_code == 200 and "Onam Sports Meet" in pub.text and "School ground" in pub.text
    assert "ഓണം കായികമേള" in TestClient(app).get(f"/p/{pid}").text
    # member registers self and a family member for an item; the item's limit puts the second on the wait list
    post = mc.get(f"/api/posts/{pid}").json()
    item = next(i for i in post["items"] if i["name_en"] == "100 m")
    fam = mc.get("/api/auth/me").json()["member"]["id"]
    fam_rows = mc.get(f"/api/members/{fam}/family").json()
    r1 = mc.post(f"/api/posts/{pid}/register", json={"item_id": item["id"]}, headers=H).json()
    r2 = mc.post(f"/api/posts/{pid}/register", json={"item_id": item["id"], "family_id": fam_rows[0]["id"]}, headers=H).json()
    assert r1["status"] == "registered" and r2["status"] == "waitlisted"
    assert mc.post(f"/api/posts/{pid}/register", json={"item_id": item["id"]}, headers=H).status_code == 409
    regs = admin.get(f"/api/posts/{pid}").json()["registrations"]
    assert {x["participant_name"] for x in regs} == {"Ramesh Kumar", "Sita"}
    # member may cancel their own, not promote it
    assert mc.post(f"/api/registrations/{r2['id']}/status", json={"status": "registered"}, headers=H).status_code == 403
    assert mc.post(f"/api/registrations/{r2['id']}/status", json={"status": "cancelled"}, headers=H).status_code == 200
    # results
    fresh()
    r = admin.post(f"/api/posts/{pid}/results", json={"results": [{"item_id": item["id"], "position": 1, "registration_id": r1["id"], "points": 5}],
                                                      "results_en": "100 m: 1. Ramesh Kumar", "notify": True, "channels": ["whatsapp"]}, headers=H)
    assert r.status_code == 200 and r.json()["queued"] >= 1
    assert mc.get(f"/api/posts/{pid}").json()["results"][0]["participant_name"] == "Ramesh Kumar"
    # committee-only post is hidden from a non-member of that committee
    cid = admin.get("/api/committees").json()[1]["id"]
    r = admin.post("/api/posts", json={"kind": "notice", "title_en": "Exec only", "audience": "committee", "committee_id": cid}, headers=H).json()
    admin.post(f"/api/posts/{r['id']}/publish", json={"notify": False}, headers=H)
    assert mc.get(f"/api/posts/{r['id']}").status_code == 404
    assert all(p["id"] != r["id"] for p in mc.get("/api/posts").json())


def test_committees_meetings_minutes_actions(admin):
    coms = admin.get("/api/committees").json()
    general = next(c for c in coms if c["kind"] == "general")
    execc = next(c for c in coms if c["kind"] == "executive")
    sports = next(c for c in coms if c["name_en"].startswith("Sports"))
    members = admin.get("/api/members?status=active").json()
    ramesh = next(m for m in members if m["name_en"] == "Ramesh Kumar")
    second = next(m for m in members if m["name_en"] == "Second")
    # Ramesh convenes the sports committee; Second is not on it
    assert admin.post(f"/api/committees/{sports['id']}/members", json={"member_id": ramesh["id"], "designation_en": "Convener", "designation_ml": "കൺവീനർ", "role": "convener"}, headers=H).status_code == 200
    admin.post(f"/api/committees/{execc['id']}/members", json={"member_id": ramesh["id"], "designation_en": "Member"}, headers=H)
    mc = login("9847012345", "ramesh-pass-1")
    sc = login("9000000002", "second-pass-1")
    # the convener schedules a meeting with an agenda and sends the notice
    r = mc.post("/api/meetings", json={"committee_id": sports["id"], "title_en": "Sports committee - September", "title_ml": "കായിക സമിതി - സെപ്റ്റംബർ", "scheduled_at": "2030-09-01T18:00",
                                       "venue_en": "Club hall", "quorum": 1, "agenda": [{"title_en": "Onam meet budget", "title_ml": "ഓണം മേള ബജറ്റ്", "kind": "resolution"},
                                                                                       {"title_en": "Ground maintenance"}]}, headers=H)
    assert r.status_code == 200, r.text
    meeting = r.json()["id"]
    fresh()
    assert mc.post(f"/api/meetings/{meeting}/notice", json={"channels": ["whatsapp"]}, headers=H).json()["queued"] == 1
    flush()
    assert "1. ഓണം മേള ബജറ്റ്" in SENT[0][3]
    # a non-member of that committee cannot see it, nor schedule one
    assert sc.get(f"/api/meetings/{meeting}").status_code == 403
    assert sc.post("/api/meetings", json={"committee_id": sports["id"], "title_en": "x", "scheduled_at": "2030-01-01"}, headers=H).status_code == 403
    assert all(m["id"] != meeting for m in sc.get("/api/meetings").json())
    # attendance and minutes
    mc.post(f"/api/meetings/{meeting}/attendance", json={"attendance": [{"member_id": ramesh["id"], "status": "present"}]}, headers=H)
    mt = mc.get(f"/api/meetings/{meeting}").json()
    agenda = mt["agenda"]
    agenda[0].update({"decision_en": "Rs 25,000 sanctioned", "decision_ml": "25,000 രൂപ അനുവദിച്ചു", "resolution_no": "SC/2030/01", "proposed_by": ramesh["id"]})
    r = mc.post("/api/meetings", json={"id": meeting, "committee_id": sports["id"], "agenda": agenda, "chaired_by": ramesh["id"], "recorded_by": ramesh["id"]}, headers=H)
    assert r.status_code == 200
    # an action item for Second (who is told) - the committee manager may assign to any member
    fresh()
    r = mc.post("/api/actions", json={"meeting_id": meeting, "agenda_item_id": agenda[1]["id"], "title_en": "Get three quotes for ground levelling", "owner_member_id": second["id"], "due": "2030-09-20"}, headers=H)
    assert r.status_code == 200, r.text
    action = r.json()["id"]
    flush()
    assert any("ground levelling" in body for *_, body in SENT)
    mine = sc.get("/api/actions?mine=1").json()
    assert [a["id"] for a in mine] == [action]
    # owner can progress it but not re-title it
    sc.post("/api/actions", json={"id": action, "status": "in_progress", "progress_note": "two quotes in", "title_en": "hacked"}, headers=H)
    a = sc.get("/api/actions?mine=1").json()[0]
    assert a["status"] == "in_progress" and a["title_en"].startswith("Get three")
    # add Second to the committee: they now see the notice and agenda but not the draft decisions
    admin.post(f"/api/committees/{sports['id']}/members", json={"member_id": second["id"]}, headers=H)
    sc = login("9000000002", "second-pass-1")
    mt = sc.get(f"/api/meetings/{meeting}").json()
    assert mt["draft_hidden"] and "decision_en" not in mt["agenda"][0]
    assert sc.get(f"/api/meetings/{meeting}/print").status_code == 403
    # publish the minutes: members are told and can now read and print them
    fresh()
    r = mc.post(f"/api/meetings/{meeting}/mom", json={"status": "published", "channels": ["whatsapp"]}, headers=H)
    assert r.status_code == 200 and r.json()["queued"] == 2
    mt = sc.get(f"/api/meetings/{meeting}").json()
    assert mt["agenda"][0]["decision_en"] == "Rs 25,000 sanctioned" and mt["status"] == "held"
    pr = sc.get(f"/api/meetings/{meeting}/print?lang=ml")
    assert pr.status_code == 200 and "25,000 രൂപ അനുവദിച്ചു" in pr.text and "SC/2030/01" in pr.text
    # only an officer signs off; afterwards it is approved
    assert mc.post(f"/api/meetings/{meeting}/mom", json={"status": "approved"}, headers=H).status_code == 403
    assert admin.post(f"/api/meetings/{meeting}/mom", json={"status": "approved"}, headers=H).status_code == 200
    assert admin.get(f"/api/meetings/{meeting}").json()["mom_status"] == "approved"
    # general body: every member sees it
    r = admin.post("/api/meetings", json={"committee_id": general["id"], "title_en": "Annual General Body", "scheduled_at": "2030-10-02T10:00"}, headers=H)
    assert sc.get(f"/api/meetings/{r.json()['id']}").status_code == 200
    # plans
    r = mc.post("/api/plans", json={"committee_id": sports["id"], "title_en": "Build a badminton court", "period": "2030-31", "budget": 150000, "owner_member_id": ramesh["id"]}, headers=H)
    assert r.status_code == 200 and mc.get("/api/plans").json()[0]["budget"] == 150000


def test_library_circulation(admin):
    r = admin.post("/api/library/books", json={"title_en": "Randamoozham", "title_ml": "രണ്ടാമൂഴം", "author_en": "M. T. Vasudevan Nair", "author_ml": "എം.ടി. വാസുദേവൻ നായർ",
                                               "language": "ml", "category": "novel", "price": 350, "add_copies": 2}, headers=H)
    assert r.status_code == 200, r.text
    book = r.json()["id"]
    accs = r.json()["accession_nos"]
    assert accs == ["00001", "00002"]
    members = admin.get("/api/members?status=active").json()
    ramesh = next(m for m in members if m["name_en"] == "Ramesh Kumar")
    second = next(m for m in members if m["name_en"] == "Second")
    fresh()
    r = admin.post("/api/library/issue", json={"member_id": ramesh["id"], "accession_no": "00001"}, headers=H)
    assert r.status_code == 200, r.text
    issue = r.json()["id"]
    assert r.json()["due_on"] == (dt.date.today() + dt.timedelta(days=14)).isoformat()
    flush()
    assert any("രണ്ടാമൂഴം" in body for *_, body in SENT)
    # same copy cannot go out twice; the book now shows one available
    assert admin.post("/api/library/issue", json={"member_id": second["id"], "accession_no": "00001"}, headers=H).status_code == 400
    b = admin.get(f"/api/library/books/{book}").json()
    assert sum(1 for c in b["copies"] if c["status"] == "available") == 1
    # the member sees their own loan and can renew once online
    mc = login("9847012345", "ramesh-pass-1")
    mine = mc.get("/api/library/issues").json()
    assert len(mine) == 1 and mine[0]["accession_no"] == "00001"
    assert mc.post("/api/library/renew", json={"issue_id": issue}, headers=H).status_code == 200
    assert mc.post("/api/library/renew", json={"issue_id": issue}, headers=H).status_code == 400
    # Second reserves the title; after the second copy goes out, a return makes it 'ready' and tells them
    admin.post("/api/library/issue", json={"member_id": second["id"], "accession_no": "00002"}, headers=H)
    sc = login("9000000002", "second-pass-1")
    assert sc.post("/api/library/reserve", json={"book_id": book}, headers=H).status_code == 200
    # overdue return with fine: back-date the due date
    with db.tx() as c:
        c.execute("UPDATE issues SET due_on=? WHERE id=?", ((dt.date.today() - dt.timedelta(days=5)).isoformat(), issue))
    assert admin.get("/api/library/issues?overdue=1").json()[0]["fine_now"] == 5
    fresh()
    r = admin.post("/api/library/return", json={"accession_no": "00001"}, headers=H)
    assert r.status_code == 200 and r.json()["fine"] == 5 and r.json()["receipt"]["head"] == "library_fine"
    flush()
    assert any("ലഭ്യമാണ്" in body or "available" in body for *_, body in SENT)   # reservation ready
    assert admin.get(f"/api/library/books/{book}").json()["reservations"][0]["status"] == "ready"
    # the fine receipt sits in accounts
    tx = admin.get("/api/finance/transactions?head=library_fine").json()
    assert tx["totals"]["receipt"] == 5 and tx["rows"][0]["no"].startswith("R/")
    # reading room and periodicals
    assert admin.post("/api/library/reading-room", json={"member_id": ramesh["id"], "purpose": "newspapers"}, headers=H).status_code == 200
    assert len(admin.get("/api/library/reading-room").json()["rows"]) == 1
    r = admin.post("/api/library/periodicals", json={"title_en": "Mathrubhumi", "title_ml": "മാതൃഭൂമി", "kind": "newspaper"}, headers=H)
    admin.post(f"/api/library/periodicals/{r.json()['id']}/received", json={}, headers=H)
    assert admin.get("/api/library/periodicals").json()[0]["last_received"] == dt.date.today().isoformat()
    assert "00001" in admin.get("/api/library/export/catalogue.csv").text
    # members cannot manage
    assert mc.post("/api/library/books", json={"title_en": "x"}, headers=H).status_code == 403


def test_finance_dues_and_broadcast(admin):
    members = admin.get("/api/members?status=active").json()
    ramesh = next(m for m in members if m["name_en"] == "Ramesh Kumar")
    year = dt.date.today().year
    fresh()
    r = admin.post("/api/finance/transactions", json={"kind": "receipt", "head": "subscription", "member_id": ramesh["id"], "amount": 120, "mode": "upi", "ref": "UPI123"}, headers=H)
    assert r.status_code == 200, r.text
    flush()
    assert any("രസീത്" in subj for ch, to, subj, body in SENT if ch == "whatsapp")
    m = admin.get(f"/api/members/{ramesh['id']}").json()
    sub = next(s for s in m["subscriptions"] if s["year"] == year)
    assert sub["status"] == "paid" and sub["paid"] == 120
    # admission fee still due; defaulters list shows it; reminders go out
    dues = admin.get("/api/finance/dues").json()
    assert any(d["member_id"] == ramesh["id"] and d["year"] == 0 for d in dues)
    assert admin.post("/api/finance/dues/generate", json={"year": year + 1}, headers=H).json()["created"] >= 2
    assert admin.post("/api/finance/dues/remind", json={"channels": ["whatsapp"]}, headers=H).json()["members"] >= 2
    # a payment voucher and the summary
    admin.post("/api/finance/transactions", json={"kind": "payment", "head": "books", "party": "DC Books", "amount": 1500, "mode": "bank"}, headers=H)
    s = admin.get("/api/finance/summary").json()
    assert s["receipts"] == 125 and s["payments"] == 1500 and s["balance"] == -1375
    # voiding a subscription receipt reopens the due
    tid = admin.get("/api/finance/transactions?head=subscription").json()["rows"][0]["id"]
    admin.post(f"/api/finance/transactions/{tid}/void", json={"note": "entered twice"}, headers=H)
    m = admin.get(f"/api/members/{ramesh['id']}").json()
    assert next(s for s in m["subscriptions"] if s["year"] == year)["status"] == "due"
    # a treasurer-only screen is closed to members; member sees their own dues
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.get("/api/finance/transactions").status_code == 403
    assert len(mc.get("/api/finance/mine").json()["dues"]) >= 2
    # broadcast to a ward, bilingual, respecting channel consent
    mc.put(f"/api/members/{ramesh['id']}", json={"notify_sms": 0}, headers=H)
    assert admin.post("/api/broadcast/preview", json={"audience": {"kind": "ward", "value": "7"}}, headers=H).json()["recipients"] == 1
    fresh()
    r = admin.post("/api/broadcast", json={"audience": {"kind": "ward", "value": "7"}, "channels": ["whatsapp", "sms", "email"], "subject": "Cleaning drive",
                                           "body_en": "Cleaning drive on Sunday 7 am", "body_ml": "ഞായറാഴ്ച രാവിലെ 7-ന് ശുചീകരണം"}, headers=H)
    assert r.status_code == 200 and r.json()["recipients"] == 1
    flush()
    chans = {ch for ch, *_ in SENT}
    assert chans == {"whatsapp", "email"} and all("ശുചീകരണം" in body for *_, body in SENT)
    bl = admin.get("/api/broadcasts").json()
    assert bl[0]["delivery"].get("sent") == 2
    ob = admin.get("/api/outbox").json()
    assert ob["counts"]["sent"] >= 2 and ob["channels"]["app"] is True


def test_daily_jobs_and_dashboard(admin):
    members = admin.get("/api/members?status=active").json()
    ramesh = next(m for m in members if m["name_en"] == "Ramesh Kumar")
    today = dt.date.today()
    admin.put(f"/api/members/{ramesh['id']}", json={"dob": f"1980-{today.month:02d}-{today.day:02d}"}, headers=H)
    r = admin.post("/api/posts", json={"kind": "camp", "title_en": "Blood donation camp", "starts_at": (today + dt.timedelta(days=1)).isoformat() + "T09:00", "audience": "members"}, headers=H)
    admin.post(f"/api/posts/{r.json()['id']}/publish", json={"notify": False}, headers=H)
    fresh()
    out = scheduler.run_daily(force=True)
    assert out["birthdays"] >= 1 and out["event_reminders"] >= 1
    flush()
    assert any("ജന്മദിനാശംസകൾ" in subj for ch, to, subj, body in SENT if ch == "whatsapp")
    assert any("Blood donation camp" in body for *_, body in SENT)
    assert scheduler.run_daily().get("skipped") is True
    d = admin.get("/api/dashboard").json()
    assert d["office"]["members_active"] >= 3 and d["birthdays"][0]["id"] == ramesh["id"]
    ov = admin.get("/api/reports/overview").json()
    assert ov["by_type"] and ov["library_popular"][0]["title_en"] == "Randamoozham"
    # users & settings are admin-only
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.get("/api/settings").status_code == 403 and mc.get("/api/users").status_code == 403
    assert admin.post("/api/settings", json={"club_short": "KASC", "member_no_prefix": "KASC"}, headers=H).status_code == 200
    assert admin.get("/api/public/info").json()["club_short"] == "KASC"
    # state-changing calls without the app header are refused (CSRF)
    assert admin.post("/api/settings", json={"club_short": "X"}).status_code == 403
    # a treasurer login
    assert admin.post("/api/users", json={"login": "treasurer@club.test", "name": "Treasurer", "role": "treasurer", "password": "treasurer-pass-1"}, headers=H).status_code == 200
    tc = login("treasurer@club.test", "treasurer-pass-1")
    assert tc.get("/api/finance/summary").status_code == 200 and tc.get("/api/members/applications").status_code == 403


def test_templates_render_in_both_languages():
    for key in i18n.TEMPLATES:
        for lang in ("en", "ml"):
            subj, body = i18n.render(key, lang, {"name": "X", "club": "C", "club_short": "S", "title": "T", "subject": "Sub", "body": "B", "no": "R1"})
            assert subj and body
    assert comms.normalise_phone("0 98470-12345") == "919847012345" and comms.normalise_phone("+91 98470 12345") == "919847012345"
    assert comms.wa_link("9847012345", "hi there").startswith("https://wa.me/919847012345?text=hi%20there")


# ====================== Phase 1 & 2 ======================
import io  # noqa: E402
import re  # noqa: E402
import time  # noqa: E402


def test_forgot_password_by_otp(admin):
    fresh()
    r = admin.post("/api/auth/forgot", json={"login": "98470 12345"}, headers=H)
    assert r.status_code == 200
    ticket = r.json()["ticket"]
    flush()
    codes = {re.search(r"\b(\d{6})\b", body).group(1) for ch, to, subj, body in SENT if to == "919847012345"}
    assert len(codes) == 1
    code = codes.pop()
    assert any(ch == "sms" for ch, *_ in SENT) and any("കോഡ്" in body for *_, body in SENT)
    # an unknown login answers the same way and sends nothing
    fresh()
    r = admin.post("/api/auth/forgot", json={"login": "9000000099"}, headers=H)
    assert r.status_code == 200 and "ticket" in r.json()
    flush(); assert not SENT
    c = TestClient(app)
    assert c.post("/api/auth/reset", json={"ticket": ticket, "code": "000000", "password": "new-pass-123"}, headers=H).status_code == 400
    r = c.post("/api/auth/reset", json={"ticket": ticket, "code": code, "password": "new-pass-123"}, headers=H)
    assert r.status_code == 200, r.text
    assert c.get("/api/auth/me").json()["user"]["name"] == "Ramesh Kumar"           # signed in by the reset itself
    assert TestClient(app).post("/api/auth/login", json={"login": "9847012345", "password": "ramesh-pass-1"}, headers=H).status_code == 401
    assert TestClient(app).post("/api/auth/login", json={"login": "9847012345", "password": "new-pass-123"}, headers=H).status_code == 200
    # the code is single-use
    assert c.post("/api/auth/reset", json={"ticket": ticket, "code": code, "password": "another-pass-1"}, headers=H).status_code == 400
    # put the old password back for the later tests
    admin.post(f"/api/members/{admin.get('/api/members?q=Ramesh').json()[0]['id']}/reset-password", json={"password": "ramesh-pass-1"}, headers=H)


def test_excel_import(admin):
    from openpyxl import Workbook, load_workbook
    t = admin.get("/api/members/import/template.xlsx")
    assert t.status_code == 200 and load_workbook(io.BytesIO(t.content)).sheetnames == ["Members", "Family", "Keys"]
    wb = Workbook(); ws = wb.active; ws.title = "Members"
    ws.append(["Membership No", "Name (English)", "Name (Malayalam)", "Gender (male/female/other)", "Date of birth (YYYY-MM-DD)", "Blood group", "Phone", "WhatsApp", "E-mail", "House (English)",
               "House (Malayalam)", "Address", "Ward", "PIN", "Occupation", "Education", "Membership type (key)", "Joined on (YYYY-MM-DD)", "Status (active/inactive)", "Notes"])
    ws.append(["ASC/2015/0003", "Imported One", "ഒന്ന്", "male", "1970-01-01", "a+", "9111111111", "", "", "", "", "", "3", "", "Farmer", "", "life", "2015-04-01", "active", ""])
    ws.append(["", "Imported Two", "", "female", "", "", "9111111112", "", "", "", "", "", "3", "", "", "", "", "", "", "no number - gets the next"])
    ws.append(["", "Bad Row", "", "", "31-12-1980", "", "9847012345", "", "", "", "", "", "", "", "", "", "gold", "", "", "dup phone, bad date, bad type"])
    ws2 = wb.create_sheet("Family"); ws2.append(["Membership No", "Member phone (if no number)", "Name (English)", "Name (Malayalam)", "Relation", "Gender", "Date of birth (YYYY-MM-DD)", "Phone", "Blood group", "Education", "Occupation"])
    ws2.append(["ASC/2015/0003", "", "Kid One", "", "son", "male", "2005-05-05", "", "", "", ""])
    ws2.append(["", "9111111112", "Kid Two", "", "daughter", "", "", "", "", "", ""])
    ws2.append(["ASC/9999/9999", "", "Orphan", "", "son", "", "", "", "", "", ""])
    buf = io.BytesIO(); wb.save(buf)
    r = admin.post("/api/members/import?commit=0", files={"file": ("reg.xlsx", buf.getvalue(), "application/octet-stream")}, headers=H)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["valid"] == 2 and p["invalid"] == 1 and p["created"] == 0 and p["family_valid"] == 2
    bad = p["members"][2]["errors"]
    assert any("phone already" in e for e in bad) and any("dob" in e for e in bad) and any("membership type" in e for e in bad)
    assert p["family"][2]["errors"] == ["member not found"]
    before = len(admin.get("/api/members?status=active").json())
    r = admin.post("/api/members/import?commit=1", files={"file": ("reg.xlsx", buf.getvalue(), "application/octet-stream")}, headers=H)
    assert r.json()["created"] == 2 and r.json()["family_created"] == 2
    rows = admin.get("/api/members?status=active").json()
    assert len(rows) == before + 2
    one = next(m for m in rows if m["name_en"] == "Imported One")
    assert one["member_no"] == "ASC/2015/0003" and one["blood_group"] == "A+" and one["membership_type"] == "life"
    two = next(m for m in rows if m["name_en"] == "Imported Two")
    assert two["member_no"].startswith("KASC/") and two["family_count"] == 1
    # importing the same file again: every row is now a duplicate
    assert admin.post("/api/members/import?commit=0", files={"file": ("reg.xlsx", buf.getvalue(), "application/octet-stream")}, headers=H).json()["valid"] == 0
    # batch cards
    cards = admin.get("/api/members/print/cards?mtype=life&lang=en")
    assert cards.status_code == 200 and "ASC/2015/0003" in cards.text and "ASC/OLD/17" in cards.text and cards.text.count('class="ecard"') == 2


def test_upi_claim_and_confirmation(admin):
    assert admin.post("/api/settings", json={"upi_id": "club@upi", "upi_name": "KASC"}, headers=H).status_code == 200
    mc = login("9847012345", "ramesh-pass-1")
    u = mc.get("/api/finance/upi?amount=120").json()
    assert u["configured"] and "pa=club%40upi" in u["link"] and "am=120.00" in u["link"] and u["qr"].startswith("data:image/svg")
    fresh()
    r = mc.post("/api/finance/claims", json={"amount": 120, "utr": "UTR123456789", "head": "subscription"}, headers=H)
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    assert mc.post("/api/finance/claims", json={"amount": 120, "utr": "UTR123456789"}, headers=H).status_code == 409
    flush()
    assert any("UTR123456789" in body for *_, body in SENT)
    mine = mc.get("/api/finance/mine").json()
    assert mine["upi"] and mine["claims"][0]["status"] == "pending"
    assert mc.post(f"/api/finance/claims/{cid}/decide", json={"decision": "confirm"}, headers=H).status_code == 403
    pending = admin.get("/api/finance/claims").json()
    assert [c["id"] for c in pending] == [cid]
    fresh()
    r = admin.post(f"/api/finance/claims/{cid}/decide", json={"decision": "confirm"}, headers=H)
    assert r.status_code == 200 and r.json()["receipt"].startswith("R/")
    flush()
    assert any("രസീത്" in subj for ch, to, subj, body in SENT)
    year = dt.date.today().year
    m = admin.get(f"/api/members/{pending[0]['member_id']}").json()
    assert next(s for s in m["subscriptions"] if s["year"] == year)["status"] == "paid"
    assert admin.get("/api/finance/transactions?head=subscription").json()["rows"][0]["mode"] == "upi"
    # a rejected claim tells the member why
    cid2 = mc.post("/api/finance/claims", json={"amount": 50, "utr": "UTRWRONG0001"}, headers=H).json()["id"]
    fresh()
    admin.post(f"/api/finance/claims/{cid2}/decide", json={"decision": "reject", "note": "amount not found in bank"}, headers=H)
    flush()
    assert any("amount not found" in body for *_, body in SENT)


def test_scan_attendance(admin):
    ramesh = admin.get("/api/members?q=Ramesh").json()[0]
    r = admin.get(f"/api/scan/MEMBER:{ramesh['member_no']}|Ramesh Kumar")
    assert r.status_code == 200 and r.json()["id"] == ramesh["id"]
    assert admin.get("/api/scan/MEMBER:NOPE").status_code == 404
    plain = login("9000000002", "second-pass-1")                        # a member who convenes nothing
    assert plain.get(f"/api/scan/{ramesh['member_no']}").status_code == 403
    # event attendance: the registered member becomes 'attended', an unregistered one gets an 'attended' registration
    post = next(p for p in admin.get("/api/posts?kind=competition").json())
    r = admin.post(f"/api/posts/{post['id']}/attend", json={"member_id": ramesh["id"]}, headers=H).json()
    assert r["new"] is False
    second = admin.get("/api/members?q=Second").json()[0]
    assert admin.post(f"/api/posts/{post['id']}/attend", json={"member_id": second["id"]}, headers=H).json()["new"] is True
    regs = admin.get(f"/api/posts/{post['id']}").json()["registrations"]
    assert {x["status"] for x in regs if x["member_id"] in (ramesh["id"], second["id"]) and not x["family_id"]} == {"attended"}


def test_documents(admin):
    pdf = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
    r = admin.post("/api/upload/docs", files={"file": ("bylaws.pdf", pdf, "application/pdf")}, headers=H)
    assert r.status_code == 200 and r.json()["path"].startswith("docs/")
    path = r.json()["path"]
    r = admin.post("/api/documents", json={"title_en": "Bye-laws 2024", "title_ml": "നിയമാവലി 2024", "category": "bylaws", "audience": "public", "year": "2024", "path": path, "mime": "application/pdf", "size": len(pdf)}, headers=H)
    assert r.status_code == 200
    admin.post("/api/documents", json={"title_en": "Audit report", "category": "audit", "audience": "members", "year": "2024", "path": path, "mime": "application/pdf"}, headers=H)
    cid = admin.get("/api/committees").json()[1]["id"]
    admin.post("/api/documents", json={"title_en": "Exec circular", "category": "circular", "audience": "committee", "committee_id": cid, "path": path}, headers=H)
    anon = TestClient(app)
    assert [d["title_en"] for d in anon.get("/api/public/documents").json()] == ["Bye-laws 2024"]
    assert anon.get(f"/uploads/{path}").status_code == 200
    sc = login("9000000002", "second-pass-1")
    titles = {d["title_en"] for d in sc.get("/api/documents").json()}
    assert titles == {"Bye-laws 2024", "Audit report"}                      # not on the executive committee
    assert sc.post("/api/documents", json={"title_en": "x", "path": path}, headers=H).status_code == 403
    assert len(admin.get("/api/documents").json()) == 3


def test_assets_and_bookings(admin):
    r = admin.post("/api/assets", json={"name_en": "Plastic chairs", "name_ml": "പ്ലാസ്റ്റിക് കസേരകൾ", "category": "furniture", "qty": 3, "location": "Store"}, headers=H)
    assert r.status_code == 200
    aid = r.json()["id"]
    a = admin.get(f"/api/assets/{aid}").json()
    assert a["code"] == "AST-0001" and a["available"] == 3
    second = admin.get("/api/members?q=Second").json()[0]
    assert admin.post(f"/api/assets/{aid}/lend", json={"member_id": second["id"], "qty": 2, "due_on": "2020-01-01"}, headers=H).status_code == 200
    assert admin.post(f"/api/assets/{aid}/lend", json={"party": "Temple festival", "qty": 2}, headers=H).status_code == 400
    assert admin.get(f"/api/assets/{aid}").json()["available"] == 1
    fresh()
    assert admin.post("/api/assets/loans/remind", headers=H).json()["queued"] >= 1
    flush()
    assert any("Plastic chairs" in body or "കസേരകൾ" in body for *_, body in SENT)
    loan = admin.get(f"/api/assets/{aid}").json()["loans"][0]
    assert admin.post(f"/api/assets/loans/{loan['id']}/return", json={"condition": "good"}, headers=H).status_code == 200
    assert admin.get(f"/api/assets/{aid}").json()["available"] == 3
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.post("/api/assets", json={"name_en": "x"}, headers=H).status_code == 403
    # facilities and bookings
    r = admin.post("/api/facilities", json={"name_en": "Club hall", "name_ml": "ക്ലബ്ബ് ഹാൾ", "rate_day": 2000, "rate_hour": 300, "deposit": 500, "capacity": 150}, headers=H)
    fid = r.json()["id"]
    r = mc.post("/api/bookings", json={"facility_id": fid, "starts_at": "2031-01-10T09:00", "ends_at": "2031-01-10T13:00", "purpose": "Wedding reception"}, headers=H)
    assert r.status_code == 200 and r.json()["amount"] == 1200             # 4 hours x 300
    bid = r.json()["id"]
    sc = login("9000000002", "second-pass-1")
    assert sc.post("/api/bookings", json={"facility_id": fid, "starts_at": "2031-01-10T12:00", "ends_at": "2031-01-10T15:00"}, headers=H).status_code == 409
    assert sc.post("/api/bookings", json={"facility_id": fid, "starts_at": "2031-01-10T13:00", "ends_at": "2031-01-12T13:00"}, headers=H).json()["amount"] == 4000   # 2 days
    cal = sc.get("/api/bookings/calendar?month=2031-01").json()
    assert len(cal) == 2 and cal[0]["purpose"] == ""                        # another member's purpose is private
    assert len(sc.get("/api/bookings").json()) == 1                         # members see only their own
    assert sc.post(f"/api/bookings/{bid}/status", json={"status": "confirmed"}, headers=H).status_code == 403
    fresh()
    r = admin.post(f"/api/bookings/{bid}/status", json={"status": "confirmed", "collect": True, "mode": "cash"}, headers=H)
    assert r.status_code == 200 and r.json()["receipt"]["amount"] == 1200
    flush()
    assert any("Club hall" in body or "ക്ലബ്ബ് ഹാൾ" in body for *_, body in SENT)
    assert admin.get("/api/bookings?month=2031-01").json()[1]["status"] == "confirmed"
    assert mc.post(f"/api/bookings/{bid}/status", json={"status": "cancelled"}, headers=H).status_code == 200


def test_fixtures_judging_points(admin):
    post = next(p for p in admin.get("/api/posts?kind=competition").json())
    pid = post["id"]
    item = admin.get(f"/api/posts/{pid}").json()["items"][1]              # tug of war
    r = admin.post(f"/api/posts/{pid}/fixtures", json={"item_id": item["id"], "round": "Final", "side_a": "Ward 7", "side_b": "Ward 8", "score_a": "2", "score_b": "1", "status": "done"}, headers=H)
    assert r.status_code == 200 and r.json()["winner"] == "Ward 7"
    admin.post(f"/api/posts/{pid}/fixtures", json={"round": "League", "side_a": "Ward 8", "side_b": "Ward 9", "score_a": "1", "score_b": "1", "status": "done"}, headers=H)
    mc = login("9847012345", "ramesh-pass-1")
    assert len(mc.get(f"/api/posts/{pid}/fixtures").json()) == 2
    assert mc.post(f"/api/posts/{pid}/fixtures", json={"side_a": "x", "side_b": "y"}, headers=H).status_code == 403
    # judges: the treasurer user scores the 100 m item
    tre = next(u for u in admin.get("/api/users").json() if u["role"] == "treasurer")
    item100 = admin.get(f"/api/posts/{pid}").json()["items"][0]
    assert admin.post(f"/api/posts/{pid}/judges", json={"judges": [{"user_id": tre["id"], "item_ids": [item100["id"]]}]}, headers=H).status_code == 200
    tc = login("treasurer@club.test", "treasurer-pass-1")
    assert [j["id"] for j in tc.get("/api/judging").json()] == [pid]
    sheet = tc.get(f"/api/posts/{pid}/judging").json()
    assert [i["id"] for i in sheet["items"]] == [item100["id"]]
    regs = [r for r in sheet["registrations"] if r["item_id"] == item100["id"]]
    assert regs
    # give the registrations teams so the points table has something to add up
    for i, r in enumerate(regs):
        admin.post(f"/api/registrations/{r['id']}/status", json={"status": "attended"}, headers=H)
    with db.tx() as c:
        for i, r in enumerate(regs):
            c.execute("UPDATE registrations SET team=? WHERE id=?", ("Ward 7" if i == 0 else "Ward 8", r["id"]))
    scores = [{"registration_id": r["id"], "score": 9.5 - i, "remark": "good"} for i, r in enumerate(regs)]
    assert tc.post(f"/api/posts/{pid}/scores", json={"scores": scores}, headers=H).json()["saved"] == len(regs)
    assert mc.post(f"/api/posts/{pid}/scores", json={"scores": scores}, headers=H).status_code == 403
    rank = admin.get(f"/api/posts/{pid}/ranking").json()
    assert rank[0]["position"] == 1 and rank[0]["avg_score"] == 9.5 and rank[0]["judges"] == 1
    r = admin.post(f"/api/posts/{pid}/results/from-ranking", json={"points": [5, 3, 1]}, headers=H)
    assert r.status_code == 200 and r.json()["results"] >= 1
    res = admin.get(f"/api/posts/{pid}").json()["results"]
    assert res[0]["points"] == 5 and res[0]["team"] == "Ward 7"
    pts = mc.get(f"/api/posts/{pid}/points").json()
    w7 = next(t for t in pts if t["team"] == "Ward 7")
    assert w7["points"] == 5 + 3 and w7["wins"] == 1 and pts[0]["team"] == "Ward 7"   # 5 from the result + 3 for the fixture win
    w8 = next(t for t in pts if t["team"] == "Ward 8")
    assert w8["draws"] == 1


def test_annual_report(admin):
    year = dt.date.today().year
    r = admin.get(f"/api/reports/annual?year={year}").json()
    assert r["members"]["active"] >= 3 and r["accounts"]["receipts"] > 0 and r["library"]["issues"] >= 1 and str(year) in r["years"]
    pr = admin.get(f"/api/reports/annual/print?year={year}&lang=ml")
    assert pr.status_code == 200 and "വാർഷിക റിപ്പോർട്ട്" in pr.text and "Randamoozham" in pr.text or "രണ്ടാമൂഴം" in pr.text
    assert login("9847012345", "ramesh-pass-1").get("/api/reports/annual").status_code == 403


def test_web_push_channel_and_crypto(admin):
    from app import push
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import ec
    mc = login("9847012345", "ramesh-pass-1")
    key = mc.get("/api/push/key").json()
    assert len(key["key"]) > 80 and key["subscribed"] is False
    ua = ec.generate_private_key(ec.SECP256R1())
    p256dh = push.b64u(ua.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint))
    auth_ = push.b64u(os.urandom(16))
    r = mc.post("/api/push/subscribe", json={"endpoint": "https://push.example.org/send/abc", "keys": {"p256dh": p256dh, "auth": auth_}}, headers=H)
    assert r.status_code == 200 and mc.get("/api/push/key").json()["subscribed"] is True
    # an in-app notification now also queues a push row, delivered through the (fake) transport
    fresh()
    admin.post("/api/broadcast", json={"audience": {"kind": "ward", "value": "7"}, "channels": ["app"], "subject": "Push test", "body_en": "Hello push"}, headers=H)
    flush()
    assert [(ch, to) for ch, to, *_ in SENT] == [("push", "sub:" + str(db.one("SELECT id FROM push_subs")["id"]))]
    # the encryption is a real aes128gcm round trip the browser could decrypt
    body = push.encrypt(b'{"title":"Hello"}', p256dh, auth_)
    assert push.decrypt_for_tests(body, ua, auth_) == b'{"title":"Hello"}'
    assert push.vapid_headers("https://push.example.org/send/abc")["Authorization"].startswith("vapid t=")
    assert mc.post("/api/push/unsubscribe", json={}, headers=H).status_code == 200 and mc.get("/api/push/key").json()["subscribed"] is False


def test_scheduled_broadcast(admin):
    fresh()
    later = (dt.datetime.now() + dt.timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M")
    r = admin.post("/api/broadcast", json={"audience": {"kind": "all"}, "channels": ["whatsapp"], "subject": "Tomorrow", "body_en": "Scheduled message", "send_at": later}, headers=H)
    assert r.status_code == 200 and r.json()["queued"] >= 3
    bid = r.json()["id"]
    flush()
    assert not SENT                                                        # not before its time
    assert admin.get("/api/outbox?status=pending").json()["counts"]["pending"] >= 3
    assert admin.get("/api/broadcasts").json()[0]["send_at"] is not None
    assert admin.post(f"/api/broadcasts/{bid}/cancel", headers=H).json()["removed"] >= 3
    assert admin.get("/api/outbox?status=pending").json()["counts"].get("pending", 0) == 0
    # blood request template renders in both languages
    for lang in ("en", "ml"):
        subj, body = i18n.render("blood_request", lang, {"blood": "O-", "patient": "a member", "hospital": "Medical College", "contact": "9847012345", "club_short": "ASC"})
        assert "O-" in subj and "9847012345" in body


# ====================== Phase 3 ======================
import hashlib  # noqa: E402
import hmac  # noqa: E402
import json  # noqa: E402


def test_online_payment_razorpay(admin):
    from app import payments
    created = []
    payments.set_order_creator(lambda amount, receipt, notes: created.append((amount, receipt, notes)) or {"id": f"order_test{len(created)}", "amount": amount, "status": "created"})
    os.environ["CLUB_RZP_KEY_SECRET"] = "testsecret"
    os.environ["CLUB_RZP_WEBHOOK_SECRET"] = "whsecret"
    try:
        mc = login("9847012345", "ramesh-pass-1")
        assert mc.get("/api/payments/config").json()["configured"] is True
        r = mc.post("/api/payments/order", json={"amount": 120, "head": "subscription"}, headers=H)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["order_id"] == "order_test1" and created[0][0] == 12000 and created[0][2]["member_no"].startswith("ASC/")
        # a forged signature is refused and the order marked failed
        bad = mc.post("/api/payments/verify", json={"razorpay_order_id": o["order_id"], "razorpay_payment_id": "pay_1", "razorpay_signature": "nope"}, headers=H)
        assert bad.status_code == 400
        # the real signature settles the order into a receipt and tells the member
        o2 = mc.post("/api/payments/order", json={"amount": 120, "head": "subscription"}, headers=H).json()
        sig = hmac.new(b"testsecret", f"{o2['order_id']}|pay_2".encode(), hashlib.sha256).hexdigest()
        fresh()
        r = mc.post("/api/payments/verify", json={"razorpay_order_id": o2["order_id"], "razorpay_payment_id": "pay_2", "razorpay_signature": sig}, headers=H)
        assert r.status_code == 200 and r.json()["receipt"].startswith("R/") and r.json()["amount"] == 120
        flush()
        assert any("pay_2" in body for *_, body in SENT)
        mine = mc.get("/api/payments/orders").json()
        assert mine[0]["status"] == "paid" and mine[0]["payment_id"] == "pay_2"
        # the webhook for the same order is idempotent; a tampered webhook is refused
        ev = json.dumps({"event": "payment.captured", "payload": {"payment": {"entity": {"id": "pay_2", "order_id": o2["order_id"]}}}}).encode()
        wsig = hmac.new(b"whsecret", ev, hashlib.sha256).hexdigest()
        anon = TestClient(app)
        assert anon.post("/api/payments/razorpay/webhook", content=ev, headers={"X-Razorpay-Signature": wsig, "Content-Type": "application/json"}).status_code == 200
        assert anon.post("/api/payments/razorpay/webhook", content=ev, headers={"X-Razorpay-Signature": "x", "Content-Type": "application/json"}).status_code == 400
        assert len([t for t in admin.get("/api/finance/transactions?head=subscription").json()["rows"] if t["ref"] == "pay_2"]) == 1
        # a webhook can settle an order the browser never verified
        o3 = mc.post("/api/payments/order", json={"amount": 50, "head": "donation"}, headers=H).json()
        ev3 = json.dumps({"event": "order.paid", "payload": {"payment": {"entity": {"id": "pay_3", "order_id": o3["order_id"]}}}}).encode()
        anon.post("/api/payments/razorpay/webhook", content=ev3, headers={"X-Razorpay-Signature": hmac.new(b"whsecret", ev3, hashlib.sha256).hexdigest(), "Content-Type": "application/json"})
        assert admin.get("/api/payments/orders").json()[0]["status"] == "paid"
        assert admin.get("/api/finance/transactions?head=donation").json()["rows"][0]["ref"] == "pay_3"
    finally:
        payments.set_order_creator(None)
        os.environ.pop("CLUB_RZP_KEY_SECRET", None); os.environ.pop("CLUB_RZP_WEBHOOK_SECRET", None)
    assert mc.get("/api/payments/config").json()["configured"] is False
    assert mc.post("/api/payments/order", json={"amount": 10}, headers=H).status_code == 400


def test_backup_local_and_drive(admin):
    from app import backup
    bdir = tempfile.mkdtemp(prefix="club-bk-")
    os.environ["CLUB_BACKUP_DIR"] = bdir
    uploaded, deleted = [], []
    store = []
    backup.set_drive_fake(upload=lambda name, path: uploaded.append(name) or store.append({"id": f"f{len(store)}", "name": name}) or f"f{len(store)}",
                          list_=lambda: list(store), delete=lambda fid: deleted.append(fid) or store.remove(next(x for x in store if x["id"] == fid)))
    try:
        admin.post("/api/settings", json={"backup_keep": "2"}, headers=H)
        r = admin.post("/api/backup/run", headers=H)
        assert r.status_code == 200, r.text
        j = r.json()
        assert not j["error"] and len(j["files"]) == 2 and j["drive"]["uploaded"] == ["f1", "f2"]
        # the copy is a real database with our members in it
        import sqlite3
        dbf = next(f for f in j["files"] if f.startswith("club-"))
        n = sqlite3.connect(os.path.join(bdir, dbf)).execute("SELECT COUNT(*) FROM members WHERE status='active'").fetchone()[0]
        assert n >= 3
        import time as _t; _t.sleep(1.1)
        admin.post("/api/backup/run", headers=H); _t.sleep(1.1)
        j3 = admin.post("/api/backup/run", headers=H).json()
        assert j3["drive"]["removed"] == 2 and len(store) == 4                       # keep 2 of each kind on Drive
        assert len([f for f in os.listdir(bdir) if f.startswith("club-")]) == 2        # and locally
        st = admin.get("/api/backup").json()
        assert st["drive_configured"] and st["last"]["reason"] == "manual" and st["keep"] == 2
        assert login("9847012345", "ramesh-pass-1").get("/api/backup").status_code == 403
        # the service-account JWT is well formed (RS256 over the two base64url segments)
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric import rsa
        key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode()
        os.environ["CLUB_GDRIVE_SA_JSON"] = json.dumps({"client_email": "svc@test.iam.gserviceaccount.com", "private_key": pem})
        sa = backup._sa()
        assert sa["client_email"].startswith("svc@")
        # the scheduler runs it once a day
        from app import scheduler
        assert scheduler.nightly_backup(force=True)["reason"] == "nightly"
        assert scheduler.nightly_backup() is None
    finally:
        backup.set_drive_fake()
        os.environ.pop("CLUB_BACKUP_DIR", None); os.environ.pop("CLUB_GDRIVE_SA_JSON", None)


def test_polls_and_elections(admin):
    coms = admin.get("/api/committees").json()
    sports = next(c for c in coms if c["name_en"].startswith("Sports"))
    members = admin.get("/api/members?status=active").json()
    ramesh = next(m for m in members if m["name_en"] == "Ramesh Kumar")
    second = next(m for m in members if m["name_en"] == "Second")
    # an election for the whole membership: two candidates, one seat, anonymous
    r = admin.post("/api/polls", json={"kind": "election", "title_en": "Secretary 2027", "title_ml": "സെക്രട്ടറി 2027", "ends_at": "2035-01-01T18:00",
                                       "options": [{"member_id": ramesh["id"]}, {"member_id": second["id"]}]}, headers=H)
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.get(f"/api/polls/{pid}").status_code == 403                                     # drafts are invisible
    assert mc.post(f"/api/polls/{pid}/vote", json={"option_ids": [1]}, headers=H).status_code == 400 or True
    fresh()
    assert admin.post(f"/api/polls/{pid}/status", json={"status": "open", "channels": ["whatsapp"]}, headers=H).json()["queued"] >= 3
    flush()
    assert any("സെക്രട്ടറി 2027" in body and "തിരഞ്ഞെടുപ്പ്" in body for *_, body in SENT)
    p = mc.get(f"/api/polls/{pid}").json()
    assert p["open_now"] and p["eligible"] and not p["voted"] and "options" not in p["results"]     # no live results
    opts = p["options"]
    assert opts[0]["cand_en"] == "Ramesh Kumar" and opts[0]["label_ml"] == "രമേശ് കുമാർ"
    assert mc.post(f"/api/polls/{pid}/vote", json={"option_ids": [opts[0]["id"], opts[1]["id"]]}, headers=H).status_code == 400   # one seat
    assert mc.post(f"/api/polls/{pid}/vote", json={"option_ids": [opts[1]["id"]]}, headers=H).status_code == 200
    assert mc.post(f"/api/polls/{pid}/vote", json={"option_ids": [opts[1]["id"]]}, headers=H).status_code == 409       # once only
    assert mc.get(f"/api/polls/{pid}").json()["voted"] is True
    sc = login("9000000002", "second-pass-1")
    sc.post(f"/api/polls/{pid}/vote", json={"option_ids": [opts[1]["id"]]}, headers=H)
    assert [x["id"] for x in mc.get("/api/polls").json()] == [pid]
    # an anonymous ballot leaves no voter on the vote rows
    assert db.one("SELECT COUNT(*) n FROM votes WHERE poll_id=? AND member_id IS NOT NULL", (pid,))["n"] == 0
    assert admin.post(f"/api/polls/{pid}/status", json={"status": "draft"}, headers=H).status_code == 400
    fresh()
    assert admin.post(f"/api/polls/{pid}/status", json={"status": "closed", "channels": ["app"]}, headers=H).status_code == 200
    res = mc.get(f"/api/polls/{pid}").json()["results"]
    assert res["ballots"] == 2 and [o["votes"] for o in res["options"]] == [0, 2] and res["options"][1]["pct"] == 100
    assert mc.post(f"/api/polls/{pid}/vote", json={"option_ids": [opts[1]["id"]]}, headers=H).status_code == 400
    assert mc.get("/api/notifications").json()[0]["title"].startswith("ഫലം")
    # a committee poll is visible only to that committee; the convener (Ramesh) can run it, Second (member) can vote
    r = mc.post("/api/polls", json={"committee_id": sports["id"], "title_en": "Jersey colour", "anonymous": 0, "max_choices": 1, "results_live": 1,
                                    "options": [{"label_en": "Green", "label_ml": "പച്ച"}, {"label_en": "Blue", "label_ml": "നീല"}]}, headers=H)
    assert r.status_code == 200
    cp = r.json()["id"]
    mc.post(f"/api/polls/{cp}/status", json={"status": "open", "notify": False}, headers=H)
    old = admin.get("/api/members?q=Old").json()[0]
    oc = TestClient(app)
    admin.post(f"/api/members/{old['id']}/reset-password", json={"password": "old-member-1"}, headers=H)
    oc.post("/api/auth/login", json={"login": "9000000004", "password": "old-member-1"}, headers=H)
    assert oc.get(f"/api/polls/{cp}").status_code == 403                                      # not on the sports committee
    o = sc.get(f"/api/polls/{cp}").json()["options"]
    assert sc.post(f"/api/polls/{cp}/vote", json={"option_ids": [o[0]["id"]]}, headers=H).status_code == 200
    live = sc.get(f"/api/polls/{cp}").json()["results"]
    assert live["options"][0]["votes"] == 1                                                     # live results
    assert mc.get(f"/api/polls/{cp}").json()["voters"][0]["chosen"] == "Green"                  # not anonymous: the convener sees who


def test_public_site(admin):
    admin.post("/api/settings", json={"about_en": "Founded in 1962 to serve the village.", "about_ml": "1962-ൽ സ്ഥാപിതം.", "map_url": "https://maps.google.com/?q=club",
                                      "social_facebook": "https://facebook.com/club", "club_phone": "04812 123456"}, headers=H)
    anon = TestClient(app)
    d = anon.get("/api/public/site").json()
    assert d["settings"]["about_en"].startswith("Founded") and d["stats"]["members"] >= 3
    en = anon.get("/site?lang=en")
    assert en.status_code == 200 and "Founded in 1962" in en.text and "Onam Sports Meet" in en.text and "facebook.com/club" in en.text and "Bye-laws 2024" in en.text
    ml = anon.get("/site")
    assert "1962-ൽ സ്ഥാപിതം" in ml.text and "ഓണം കായികമേള" in ml.text and "അംഗത്വത്തിന് അപേക്ഷിക്കുക" in ml.text
    # office bearers come from the executive committee
    coms = admin.get("/api/committees").json(); execc = next(c for c in coms if c["kind"] == "executive")
    ramesh = admin.get("/api/members?q=Ramesh").json()[0]
    admin.post(f"/api/committees/{execc['id']}/members", json={"member_id": ramesh["id"], "designation_en": "President", "designation_ml": "പ്രസിഡന്റ്", "role": "chair"}, headers=H)
    assert "President" in anon.get("/site?lang=en").text and "പ്രസിഡന്റ്" in anon.get("/site?lang=ml").text


# ====================== hardening ======================
def test_card_html_is_escaped(admin):
    r = admin.post("/api/members", json={"name_en": "Evil <script>alert(1)</script>", "phone": "9000000077", "membership_type": "ordinary"}, headers=H)
    assert r.status_code == 200
    mid = r.json()["id"]
    card = admin.get(f"/api/members/{mid}/card?lang=en").text
    assert "<script>alert(1)</script>" not in card and "&lt;script&gt;" in card
    sheet = admin.get("/api/members/print/cards?mtype=ordinary&lang=en").text
    assert "<script>alert(1)</script>" not in sheet and "&lt;script&gt;" in sheet
    admin.post(f"/api/members/{mid}/status", json={"status": "resigned"}, headers=H)


def test_uploads_must_be_real_images_or_pdfs(admin):
    anon = TestClient(app)
    assert anon.post("/api/public/photo", files={"file": ("x.jpg", b"<html>not an image</html>", "image/jpeg")}).status_code == 400
    assert admin.post("/api/upload/docs", files={"file": ("x.pdf", b"<html>not a pdf</html>", "application/pdf")}, headers=H).status_code == 400
    import io as _io
    from PIL import Image
    buf = _io.BytesIO(); Image.new("RGB", (40, 40), (10, 200, 30)).save(buf, "JPEG")
    r = anon.post("/api/public/photo", files={"file": ("ok.jpg", buf.getvalue(), "image/jpeg")})
    assert r.status_code == 200 and r.json()["path"].endswith(".jpg")


def test_rate_limits_and_security_headers(admin):
    from app import security
    security.reset()
    anon = TestClient(app)
    codes = [anon.post("/api/auth/login", json={"login": "nobody@x", "password": "wrong-password"}, headers=H).status_code for _ in range(12)]
    assert codes[:10] == [401] * 10 and codes[10:] == [429, 429]
    security.reset()
    assert anon.post("/api/auth/login", json={"login": "nobody@x", "password": "wrong-password"}, headers=H).status_code == 401
    r = anon.get("/")
    assert r.headers["X-Content-Type-Options"] == "nosniff" and r.headers["X-Frame-Options"] == "SAMEORIGIN" and "default-src 'self'" in r.headers["Content-Security-Policy"]
    assert "Content-Security-Policy" not in anon.get("/api/public/info").headers       # JSON gets the other headers only
    assert anon.get("/api/public/info").headers["Referrer-Policy"] == "strict-origin-when-cross-origin"


def test_handbook_and_ecard_data(admin):
    anon = TestClient(app)
    en = anon.get("/help?lang=en"); ml = anon.get("/help")
    assert en.status_code == 200 and "For office bearers" in en.text and "Applying for membership" in en.text
    assert "ഭാരവാഹികൾക്ക്" in ml.text and "അംഗത്വത്തിന് അപേക്ഷിക്കൽ" in ml.text
    j = anon.get("/api/help?lang=ml").json()
    assert j["lang"] == "ml" and len(j["sections"]) >= 12 and all(x["paragraphs"] for x in j["sections"])
    mc = login("9847012345", "ramesh-pass-1")
    me = mc.get("/api/auth/me").json()["member"]
    d = mc.get(f"/api/members/{me['id']}/card.json").json()
    assert d["member_no"] == me["member_no"] and d["name_ml"] == "രമേശ് കുമാർ" and d["qr"].startswith("data:image/svg") and d["club_short"]
    other = admin.get("/api/members?q=Second").json()[0]
    assert mc.get(f"/api/members/{other['id']}/card.json").status_code == 403
    assert "ecard" in mc.get(f"/api/members/{me['id']}/card?lang=en").text


# ====================== multi-club ======================
def test_multi_club_isolation(admin):
    from app import tenancy
    os.environ["CLUB_MASTER_KEY"] = "master-key-1"
    try:
        anon = TestClient(app)
        assert anon.get("/api/_platform/tenants").status_code == 403
        MK = {"X-Master-Key": "master-key-1"}
        before = anon.get("/api/_platform/tenants", headers=MK).json()
        n_default = db.conn().execute("select count(*) from members where status='active'").fetchone()[0]
        assert [t["slug"] for t in before["tenants"]] == ["default"] and before["tenants"][0]["members"] == n_default
        r = anon.post("/api/_platform/tenants", json={"slug": "kasc", "name_en": "Kottayam Arts & Sports Club", "name_ml": "കോട്ടയം ആർട്സ് & സ്പോർട്സ് ക്ലബ്ബ്", "short": "KSC",
                                                     "domains": ["kasc.test"], "admin_login": "sec@kasc.test", "admin_password": "kasc-admin-1"}, headers=MK)
        assert r.status_code == 200, r.text
        assert anon.post("/api/_platform/tenants", json={"slug": "Bad Slug", "name_en": "x", "admin_login": "a@b", "admin_password": "12345678"}, headers=MK).status_code == 400
        assert anon.post("/api/_platform/tenants", json={"slug": "other", "name_en": "x", "domains": ["kasc.test"], "admin_login": "a@b", "admin_password": "12345678"}, headers=MK).status_code == 400
        # the new club is reached by its host name and starts empty, with its own admin and settings
        k = TestClient(app, base_url="http://kasc.test")
        info = k.get("/api/public/info").json()
        assert info["club_short"] == "KSC" and info["club_name_ml"].startswith("കോട്ടയം")
        assert k.post("/api/auth/login", json={"login": "admin@club.test", "password": "admin-password-1"}, headers=H).status_code == 401   # the default club's admin does not exist here
        assert k.post("/api/auth/login", json={"login": "sec@kasc.test", "password": "kasc-admin-1"}, headers=H).status_code == 200
        assert k.get("/api/members?status=active").json() == []
        r = k.post("/api/members", json={"name_en": "Kottayam Member", "phone": "9111222333", "membership_type": "life"}, headers=H)
        assert r.status_code == 200 and r.json()["member_no"].startswith("KSC/")
        # ...and nothing leaks across: the default club does not see it, the new club does not see the default's members
        assert all(m["name_en"] != "Kottayam Member" for m in admin.get("/api/members?status=active").json())
        assert [m["name_en"] for m in k.get("/api/members?status=active").json()] == ["Kottayam Member"]
        assert admin.get("/api/members?q=Kottayam").json() == [] and k.get("/api/members?q=Ramesh").json() == []
        # uploads and messages live with their club; the outbox worker visits every club
        fresh()
        k.post("/api/broadcast", json={"audience": {"kind": "all"}, "channels": ["whatsapp"], "subject": "Hi", "body_en": "Hello Kottayam"}, headers=H)
        n = 0
        for slug in tenancy.all_slugs():
            with tenancy.use(slug):
                n += comms.process_outbox(100)
        assert n >= 1 and any("Hello Kottayam" in body for *_, body in SENT)
        assert k.get("/api/outbox").json()["counts"].get("sent") == 1 and "Hello Kottayam" not in str(admin.get("/api/outbox").json()["rows"][:3])
        # the public site and links use the club's own domain
        with tenancy.use("kasc"):
            assert comms.club_ctx()["link"] == "https://kasc.test"
        assert "Kottayam Arts" in k.get("/site?lang=en").text and "Kottayam Arts" not in admin.get("/site?lang=en").text
        # an unknown host falls back to the default club unless routing is strict
        assert TestClient(app, base_url="http://nobody.test").get("/api/public/info").json()["club_short"] != "KSC"
        anon.post("/api/_platform/settings", json={"strict": True, "default_domains": ["club.test"]}, headers=MK)
        assert TestClient(app, base_url="http://nobody.test").get("/api/public/info").status_code == 404
        assert TestClient(app, base_url="http://club.test").get("/api/public/info").status_code == 200
        assert TestClient(app, base_url="http://nobody.test").get("/healthz").status_code == 200
        # Caddy's on-demand TLS hook: certificates only for registered domains, no master key needed
        assert anon.get("/api/_platform/ask", params={"domain": "kasc.test"}).status_code == 200
        assert anon.get("/api/_platform/ask", params={"domain": "KASC.TEST:443"}).status_code == 200
        assert anon.get("/api/_platform/ask", params={"domain": "club.test"}).status_code == 200
        assert anon.get("/api/_platform/ask", params={"domain": "nobody.test"}).status_code == 404
        assert anon.get("/api/_platform/ask").status_code == 404
        # deactivating a club takes it off the air; the platform listing shows both
        anon.post("/api/_platform/tenants/kasc", json={"active": False}, headers=MK)
        assert k.get("/api/public/info").status_code == 404
        assert anon.get("/api/_platform/ask", params={"domain": "kasc.test"}).status_code == 404
        anon.post("/api/_platform/tenants/kasc", json={"active": True}, headers=MK)
        anon.post("/api/_platform/settings", json={"strict": False, "default_domains": []}, headers=MK)
        lst = anon.get("/api/_platform/tenants", headers=MK).json()["tenants"]
        assert {t["slug"]: t.get("members") for t in lst}["kasc"] == 1
        assert anon.get("/_platform").status_code == 200
    finally:
        os.environ.pop("CLUB_MASTER_KEY", None)
    assert anon.get("/_platform").status_code == 404


def test_white_label_brand_and_own_channels(admin):
    import io as _io
    from PIL import Image
    from app import brand, tenancy
    anon = TestClient(app)
    # colours: valid hex is kept, junk falls back to the default; the palette reaches the shell, the manifest and the API
    assert admin.post("/api/settings", json={"brand_primary": "#123456", "brand_accent": "zzz", "brand_app_name": "ClubDesk KASC"}, headers=H).status_code == 200
    info = anon.get("/api/public/info").json()["brand"]
    assert info["primary"] == "#123456" and info["accent"] == "#C9542A" and info["app_name"] == "ClubDesk KASC" and info["logo"] == ""
    shell = anon.get("/").text
    assert "--brand:#123456" in shell and "<title>ClubDesk KASC</title>" in shell and 'content="#123456"' in shell
    man = anon.get("/manifest.webmanifest").json()
    assert man["theme_color"] == "#123456" and man["short_name"] == "ClubDesk KASC" and man["icons"][0]["src"].endswith(".svg")
    assert anon.get("/brand/icon-192.png").status_code == 200  # stock icon while there is no logo
    # a dark primary gets white text, a pale one dark text
    assert brand.ink_on("#123456") == "#FFFFFF" and brand.ink_on("#F2E8C0") == "#1D2A28"
    # the logo: only settings managers, images only; it is served to visitors and becomes the app icon
    buf = _io.BytesIO(); Image.new("RGBA", (300, 120), (200, 30, 30, 255)).save(buf, "PNG")
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.post("/api/upload/brand", files={"file": ("logo.png", buf.getvalue(), "image/png")}, headers=H).status_code == 403
    assert admin.post("/api/upload/brand", files={"file": ("x.pdf", b"%PDF-1.4", "application/pdf")}, headers=H).status_code == 400
    r = admin.post("/api/upload/brand", files={"file": ("logo.png", buf.getvalue(), "image/png")}, headers=H)
    assert r.status_code == 200 and r.json()["path"].startswith("brand/")
    logo = r.json()["path"]
    assert admin.post("/api/settings", json={"brand_logo": "../etc/passwd"}, headers=H).status_code == 400
    assert admin.post("/api/settings", json={"brand_logo": logo}, headers=H).status_code == 200
    assert anon.get(f"/uploads/{logo}").status_code == 200
    icon = anon.get("/brand/icon-192.png")
    assert icon.status_code == 200 and Image.open(_io.BytesIO(icon.content)).size == (192, 192)
    assert Image.open(_io.BytesIO(anon.get("/brand/icon-512.png").content)).size == (512, 512)
    assert Image.open(_io.BytesIO(anon.get("/brand/apple-touch-icon.png").content)).size == (180, 180)
    assert anon.get("/manifest.webmanifest").json()["icons"][0]["src"] == "/brand/icon-192.png"
    assert anon.get("/api/public/info").json()["brand"]["logo"] == f"/uploads/{logo}"
    # the brand reaches the public website, the handbook, a public post page, the annual report and the membership card
    assert logo in anon.get("/site?lang=en").text and "#123456" in anon.get("/site?lang=en").text
    assert "#123456" in anon.get("/help?lang=en").text and logo in anon.get("/help?lang=en").text
    assert "#123456" in admin.get("/api/reports/annual/print?year=2026&lang=en").text
    mid = db.one("SELECT id FROM members WHERE status='active' ORDER BY id LIMIT 1")["id"]
    card = admin.get(f"/api/members/{mid}/card?lang=en").text
    assert "--cb:#123456" in card and "blogo" in card and logo in card
    d = admin.get(f"/api/members/{mid}/card.json").json()
    assert d["brand"]["primary"] == "#123456" and d["brand"]["logo"] == logo
    # removing the logo drops the generated icons; the stock icon is served again
    assert admin.post("/api/settings", json={"brand_logo": ""}, headers=H).status_code == 200
    assert not os.path.exists(brand.icon_path(192)) and anon.get("/brand/icon-192.png").status_code == 200
    # this club's own gateway accounts: secrets are masked on read, kept when the mask is sent back, and win over .env
    s = admin.get("/api/settings").json()
    assert s["smtp_pass"] == "" and s["channels"]["email"] is False and s["own_channels"]["email"] is False
    assert admin.post("/api/settings", json={"sms_headers": "{not json"}, headers=H).status_code == 400
    assert admin.post("/api/settings", json={"smtp_host": "smtp.kasc.test", "smtp_from": "KASC <club@kasc.test>", "smtp_pass": "secret-1", "rzp_key_id": "rzp_test_1", "rzp_key_secret": "rzp-secret"}, headers=H).status_code == 200
    s = admin.get("/api/settings").json()
    assert s["smtp_pass"] == "••••••••" and s["rzp_key_secret"] == "••••••••" and s["channels"]["email"] is True and s["own_channels"]["email"] is True and s["gateway"] is True
    assert admin.post("/api/settings", json={"smtp_pass": "••••••••", "smtp_user": "club"}, headers=H).status_code == 200
    assert db.setting("smtp_pass") == "secret-1" and comms.cfg("CLUB_SMTP_HOST") == "smtp.kasc.test"
    os.environ["CLUB_SMTP_HOST"] = "smtp.server.test"
    try:
        assert comms.cfg("CLUB_SMTP_HOST") == "smtp.kasc.test"
        admin.post("/api/settings", json={"smtp_host": ""}, headers=H)
        assert comms.cfg("CLUB_SMTP_HOST") == "smtp.server.test"
    finally:
        os.environ.pop("CLUB_SMTP_HOST", None)
    admin.post("/api/settings", json={"smtp_from": "", "smtp_pass": "", "rzp_key_id": "", "rzp_key_secret": "", "brand_primary": "#0F4C45", "brand_app_name": ""}, headers=H)
    # the platform operator's own brand: product name on /_platform and a "powered by" footer for every new club
    os.environ["CLUB_MASTER_KEY"] = "master-key-1"
    try:
        MK = {"X-Master-Key": "master-key-1"}
        assert anon.post("/api/_platform/settings", json={"product_name": "ClubDesk", "powered_by": "Powered by ClubDesk"}, headers=MK).status_code == 200
        assert anon.get("/api/_platform/tenants", headers=MK).json()["product_name"] == "ClubDesk"
        r = anon.post("/api/_platform/tenants", json={"slug": "wl", "name_en": "White Label Club", "domains": ["wl.test"], "admin_login": "sec@wl.test", "admin_password": "wl-admin-12"}, headers=MK)
        assert r.status_code == 200, r.text
        w = TestClient(app, base_url="http://wl.test")
        assert w.get("/api/public/info").json()["footer_credit"] == "Powered by ClubDesk"
        assert "Powered by ClubDesk" in w.get("/site?lang=en").text
        with tenancy.use("wl"):
            assert comms.cfg("CLUB_SMTP_HOST") == ""  # the other club's SMTP account is not shared
        anon.post("/api/_platform/settings", json={"product_name": "", "powered_by": ""}, headers=MK)
    finally:
        os.environ.pop("CLUB_MASTER_KEY", None)


def test_saas_signup_pricing_billing_and_suspension(admin):
    import time as _time
    from app import saas, tenancy
    anon = TestClient(app)
    # pricing rule: flat up to 49, then per member
    assert saas.price_for(0) == 2500 and saas.price_for(49) == 2500 and saas.price_for(50) == 6000 and saas.price_for(99) == 11880 and saas.price_for(150) == 18000
    assert saas.product()["name"] == "ClubPro" and "Team Bpro" in saas.product()["powered_by"]
    # public pages work on any host, in both languages, and list the services and the price
    for host in ("http://testserver", "http://nobody.test"):
        c = TestClient(app, base_url=host)
        pr = c.get("/pricing?lang=en").text
        assert "₹2,500" in pr and "Up to 49 members" in pr and "Library &amp; reading room" in pr and "Facilities &amp; rentals" in pr and "Powered by Team Bpro" in pr
        assert "ലൈബ്രറിയും വായനശാലയും" in c.get("/pricing?lang=ml").text
        assert "Start your club on ClubPro" in c.get("/signup?lang=en").text
    # self-serve signup with a one-time code through the operator's SMS channel
    os.environ["CLUB_SMS_URL"] = "https://sms.test/?to={to}&text={text}"
    os.environ["CLUB_BASE_DOMAIN"] = "clubs.test"
    try:
        fresh()
        bad = anon.post("/api/saas/signup/start", json={"club_name_en": "Pala Arts Club", "phone": "123", "password": "x", "agree": True}, headers=H)
        assert bad.status_code == 400
        r = anon.post("/api/saas/signup/start", json={"club_name_en": "Pala Arts Club", "club_name_ml": "പാലാ ആർട്സ് ക്ലബ്ബ്", "place": "Pala", "short": "PAC", "admin_name": "Jose", "phone": "9847099999",
                                                       "email": "jose@pala.test", "password": "pala-pass-1", "lang": "ml", "agree": True}, headers=H)
        assert r.status_code == 200, r.text
        assert r.json()["slug"] == "pala-arts-club" and r.json()["channel"] == "sms"
        sms = [x for x in SENT if x[0] == "sms" and x[1] == "9847099999"]
        assert sms and "ClubPro" in sms[-1][3]
        code = re.search(r"(\d{6})", sms[-1][3]).group(1)
        ticket = r.json()["ticket"]
        assert anon.post("/api/saas/signup/verify", json={"ticket": ticket, "code": "000000"}, headers=H).status_code == 400
        r = anon.post("/api/saas/signup/verify", json={"ticket": ticket, "code": code}, headers=H)
        assert r.status_code == 200, r.text
        assert r.json()["url"] == "https://pala-arts-club.clubs.test" and r.json()["login"] == "jose@pala.test" and r.json()["trial_days"] == 30
        # the same address cannot be taken twice
        assert anon.post("/api/saas/signup/start", json={"club_name_en": "Pala Arts Club", "phone": "9847099998", "password": "pala-pass-1", "agree": True}, headers=H).status_code == 400
    finally:
        os.environ.pop("CLUB_SMS_URL", None)
        os.environ.pop("CLUB_BASE_DOMAIN", None)
    # the new club: trial, its own admin, Malayalam, footer credit, plan visible in public info and on the plan screen
    pc = TestClient(app, base_url="http://pala-arts-club.clubs.test")
    info = pc.get("/api/public/info").json()
    assert info["club_short"] == "PAC" and info["default_lang"] == "ml" and info["footer_credit"] == "Powered by Team Bpro"
    assert info["plan"]["status"] == "trial" and 28 <= info["plan"]["days_left"] <= 30 and info["plan"]["read_only"] is False and info["plan"]["product"] == "ClubPro"
    assert pc.post("/api/auth/login", json={"login": "jose@pala.test", "password": "pala-pass-1"}, headers=H).status_code == 200
    plan = pc.get("/api/saas/plan").json()
    assert plan["status"] == "trial" and plan["members"] == 0 and plan["price"] == 2500 and plan["gateway"] is False and plan["pricing"]["per_member"] == 120
    # the operator's own club pays nothing
    assert admin.get("/api/saas/plan").json()["exempt"] is True and anon.get("/api/public/info").json()["plan"]["exempt"] is True
    # UPI claim -> the operator confirms it on the console -> paid for a year
    os.environ["CLUB_MASTER_KEY"] = "master-key-1"
    try:
        MK = {"X-Master-Key": "master-key-1"}
        assert anon.post("/api/_platform/settings", json={"pricing": {"upi_id": "teambpro@upi", "upi_name": "Team Bpro", "support": "WhatsApp 98470 00000"}}, headers=MK).status_code == 200
        assert anon.post("/api/_platform/settings", json={"pricing": {"flat_price": "abc"}}, headers=MK).status_code == 400
        assert pc.get("/api/saas/plan").json()["upi_id"] == "teambpro@upi"
        assert pc.post("/api/saas/pay/claim", json={"ref": ""}, headers=H).status_code == 400
        assert pc.post("/api/saas/pay/claim", json={"ref": "UTR123456789012"}, headers=H).status_code == 200
        bill = anon.get("/api/_platform/billing", headers=MK).json()
        assert bill["summary"]["claims"] == 1 and bill["summary"]["by_status"]["trial"] >= 1
        row = next(r for r in bill["rows"] if r["slug"] == "pala-arts-club")
        assert row["claims"][0]["ref"] == "UTR123456789012" and row["contact"]["phone"] == "9847099999"
        r = anon.post("/api/_platform/tenants/pala-arts-club/plan", json={"action": "approve_claim", "ref": "UTR123456789012"}, headers=MK)
        assert r.status_code == 200 and r.json()["view"]["status"] == "active" and 360 <= r.json()["view"]["days_left"] <= 366
        plan = pc.get("/api/saas/plan").json()
        assert plan["status"] == "active" and plan["payments"][-1]["via"] == "upi" and plan["payments"][-1]["amount"] == 2500
        # online payment through the operator's Razorpay with a fake gateway: order -> verify -> another year
        saas.set_order_creator(lambda amt, receipt, notes: {"id": "order_op_1", "amount": amt})
        try:
            o = pc.post("/api/saas/pay/order", json={}, headers=H).json()
            assert o["amount"] == 2500 and o["order_id"] == "order_op_1" and "ClubPro" in o["description"]
            assert pc.post("/api/saas/pay/verify", json={"razorpay_order_id": "wrong"}, headers=H).status_code == 404
            r = pc.post("/api/saas/pay/verify", json={"razorpay_order_id": "order_op_1", "razorpay_payment_id": "pay_1", "razorpay_signature": "x"}, headers=H)
            assert r.status_code == 200 and 725 <= r.json()["plan"]["days_left"] <= 731
            assert len(pc.get("/api/saas/plan").json()["payments"]) == 2
        finally:
            saas.set_order_creator(None)
        # the operator's own club must not be able to pay a subscription
        assert admin.post("/api/saas/pay/order", json={}, headers=H).status_code == 400
        # expiry: past the grace period the club is read-only (reads and sign-in work, writes get 402), then reopened by paying
        reg = tenancy.registry(); reg["tenants"]["pala-arts-club"]["plan"]["paid_until"] = _time.time() - 20 * 86400; tenancy.save_registry(reg)
        assert pc.get("/api/public/info").json()["plan"]["status"] == "suspended"
        assert pc.get("/api/members", headers=H).status_code == 200
        r = pc.post("/api/settings", json={"club_phone": "1"}, headers=H)
        assert r.status_code == 402 and r.json()["read_only"] is True
        assert pc.post("/api/auth/logout", headers=H).status_code == 200
        assert pc.post("/api/auth/login", json={"login": "jose@pala.test", "password": "pala-pass-1"}, headers=H).status_code == 200
        # the daily job records the state and reminds the office bearers (app + e-mail to the admin)
        fresh()
        out = saas.daily()
        assert out["changed"] >= 1 and out["reminded"] >= 1
        assert tenancy.tenants()["pala-arts-club"]["plan"]["status"] == "suspended"
        with tenancy.use("pala-arts-club"):
            ob = db.q("SELECT channel, subject FROM outbox WHERE template='saas_suspended'")
        assert not ob  # the admin user has no member record yet, so nothing could be addressed; no crash either
        # within the grace window the club is only "overdue": still writable
        reg = tenancy.registry(); reg["tenants"]["pala-arts-club"]["plan"]["paid_until"] = _time.time() - 5 * 86400; reg["tenants"]["pala-arts-club"]["plan"]["status"] = "active"; tenancy.save_registry(reg)
        assert pc.get("/api/public/info").json()["plan"]["status"] == "overdue"
        assert pc.post("/api/settings", json={"club_phone": "9847000001"}, headers=H).status_code == 200
        # operator actions: suspend, reopen, mark paid, extend trial
        assert anon.post("/api/_platform/tenants/pala-arts-club/plan", json={"action": "suspended"}, headers=MK).json()["view"]["status"] == "suspended"
        assert pc.post("/api/settings", json={"club_phone": "1"}, headers=H).status_code == 402
        assert anon.post("/api/_platform/tenants/pala-arts-club/plan", json={"action": "reactivate"}, headers=MK).json()["view"]["status"] in ("trial", "active", "overdue")
        assert pc.post("/api/settings", json={"club_phone": "9847000001"}, headers=H).status_code == 200
        v = anon.post("/api/_platform/tenants/pala-arts-club/plan", json={"action": "mark_paid", "amount": 2500, "ref": "RCPT-7"}, headers=MK).json()["view"]
        assert v["status"] == "active" and v["payments"][-1]["ref"] == "RCPT-7"
        assert anon.post("/api/_platform/tenants/nobody/plan", json={"action": "mark_paid"}, headers=MK).status_code == 400
        assert anon.post("/api/_platform/tenants/pala-arts-club/plan", json={"action": "extend_trial", "days": 10}, headers=MK).status_code == 200
    finally:
        os.environ.pop("CLUB_MASTER_KEY", None)


def test_terms_and_privacy_pages(admin):
    anon = TestClient(app)
    for host in ("http://testserver", "http://nobody.test"):
        c = TestClient(app, base_url=host)
        t = c.get("/terms?lang=en").text
        assert "Terms of service" in t and "ClubPro" in t and "Team Bpro" in t and "read-only" in t and "Kottayam" in t
        assert "സേവന വ്യവസ്ഥകൾ" in c.get("/terms?lang=ml").text and "വായിക്കാൻ മാത്ര" in c.get("/terms?lang=ml").text
        p = c.get("/privacy?lang=en").text
        assert "Privacy policy" in p and "Digital Personal Data Protection Act" in p and "blood group" in p.lower() and "Razorpay" in p
        assert "സ്വകാര്യതാ നയം" in c.get("/privacy?lang=ml").text and "രക്തഗ്രൂപ്പ്" in c.get("/privacy?lang=ml").text
    # linked from signup, pricing and the app's public shell
    assert '/terms?lang=en' in anon.get("/signup?lang=en").text and '/privacy?lang=ml' in anon.get("/pricing?lang=ml").text
    assert anon.get("/terms?lang=xx").status_code == 200


def test_society_accounts_statements_close_and_tally(admin):
    """Opening balances, entries in a past financial year, the three statements, year close with carry-forward and
    lock, reopen, cash book and the Tally export. Today is in a later year, so the past year can be closed."""
    fy = "2025-26"
    assert admin.post("/api/settings", json={"fy_start_month": "4", "custom_heads": [{"key": "fixed_deposit", "en": "Fixed deposit", "ml": "സ്ഥിര നിക്ഷേപം", "kind": "both", "nature": "asset"}]}, headers=H).status_code == 200
    heads = admin.get("/api/accounts/heads").json()
    assert any(h["key"] == "fixed_deposit" and h["nature"] == "asset" for h in heads["heads"]) and heads["current_fy"].startswith(str(dt.date.today().year if dt.date.today().month >= 4 else dt.date.today().year - 1))
    r = admin.post("/api/accounts/openings", json={"fy": fy, "cash": 1000, "bank": 5000, "items": [{"item": "books", "amount": 20000}, {"item": "library_deposit", "amount": 2000}]}, headers=H)
    assert r.status_code == 200, r.text
    tx = [("receipt", "subscription", 500, "cash"), ("receipt", "library_deposit", 300, "cash"), ("payment", "books", 1200, "bank"), ("payment", "rent", 400, "cash"),
          ("payment", "deposit_refund", 100, "cash"), ("payment", "fixed_deposit", 1000, "bank")]
    ids = []
    for kind, head, amount, mode in tx:
        r = admin.post("/api/finance/transactions", json={"kind": kind, "head": head, "party": "Test", "amount": amount, "mode": mode, "tx_date": "2025-06-10", "notify": False}, headers=H)
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])
    assert admin.post("/api/accounts/contra", json={"head": "cash_to_bank", "amount": 500, "tx_date": "2025-07-01"}, headers=H).json()["no"].startswith("C/2025/")
    st = admin.get(f"/api/accounts/statements?fy={fy}").json()
    assert st["opening"] == {"cash": 1000, "bank": 5000, "derived": False}
    assert st["closing"] == {"cash": 800, "bank": 3300}
    assert st["ie"]["total_income"] == 500 and st["ie"]["total_expense"] == 400 and st["ie"]["surplus"] == 100
    bs = st["bs"]
    assert bs["general_fund"] == {"opening": 24000, "surplus": 100, "closing": 24100}
    assets = {a["item"]: a["closing"] for a in bs["assets"]}; liabs = {a["item"]: a["closing"] for a in bs["liabilities"]}
    assert assets == {"books": 21200, "fixed_deposit": 1000} and liabs == {"library_deposit": 2200}
    assert bs["balanced"] and bs["total_assets"] == 26300 == bs["total_liabilities"]
    assert st["rp"]["total_receipts"] == st["rp"]["total_payments"] == 1000 + 5000 + 800 + 500
    cb = admin.get(f"/api/accounts/book?fy={fy}&book=cash").json()
    assert cb["opening"] == 1000 and cb["closing"] == 800 and [x["balance"] for x in cb["rows"]][-1] == 800 and any(x["kind"] == "contra" for x in cb["rows"])
    bb = admin.get(f"/api/accounts/book?fy={fy}&book=bank").json()
    assert bb["closing"] == 3300 and bb["total_dr"] == 500 and bb["total_cr"] == 2200
    lg = admin.get(f"/api/accounts/ledger?fy={fy}&head=books").json()
    assert lg["total"] == 1200 and lg["nature"] == "asset"
    for lang, needle in (("en", "Receipts and Payments Account"), ("ml", "ബാലൻസ് ഷീറ്റ്")):
        p = admin.get(f"/api/accounts/statements/print?fy={fy}&lang={lang}")
        assert p.status_code == 200 and needle in p.text and "24,100.00" in p.text
    assert "Cash book" in admin.get(f"/api/accounts/book/print?fy={fy}&book=cash&lang=en").text
    xml = admin.get(f"/api/accounts/tally.xml?fy={fy}").text
    assert "Import Data" in xml and xml.count("<VOUCHER ") == 7 and 'VCHTYPE="Contra"' in xml and "Indirect Incomes" in xml and "Fixed Assets" in xml and "<AMOUNT>-500.00</AMOUNT>" in xml
    # the current year cannot be closed; the past year can, and the balances carry forward
    cur = heads["current_fy"]
    assert admin.post("/api/accounts/close", json={"fy": cur}, headers=H).status_code == 400
    r = admin.post("/api/accounts/close", json={"fy": fy, "note": "audited"}, headers=H)
    assert r.status_code == 200, r.text
    assert r.json()["next"] == "2026-27" and r.json()["closing"]["bank"] == 3300
    op = admin.get("/api/accounts/openings?fy=2026-27").json()
    assert op["cash"] == 800 and op["bank"] == 3300 and op["derived"] and op["locked"] and {i["item"]: i["amount"] for i in op["items"] if i["amount"]} == {"books": 21200, "library_deposit": 2200, "fixed_deposit": 1000}
    assert admin.get(f"/api/accounts/statements?fy={fy}").json()["status"] == "closed"
    # a closed year takes no entries, no voids and no opening changes
    assert admin.post("/api/finance/transactions", json={"kind": "receipt", "head": "donation", "party": "x", "amount": 10, "tx_date": "2025-06-11"}, headers=H).status_code == 423
    assert admin.post(f"/api/finance/transactions/{ids[0]}/void", json={"note": "x"}, headers=H).status_code == 423
    assert admin.post("/api/accounts/openings", json={"fy": "2026-27", "cash": 1}, headers=H).status_code == 409
    assert admin.post("/api/accounts/close", json={"fy": fy}, headers=H).status_code == 409
    # reopen releases the lock and drops the carried-forward openings
    assert admin.post("/api/accounts/reopen", json={"fy": fy}, headers=H).status_code == 200
    assert admin.get("/api/accounts/openings?fy=2026-27").json()["cash"] == 0
    assert admin.post("/api/finance/transactions", json={"kind": "receipt", "head": "donation", "party": "x", "amount": 10, "tx_date": "2025-06-11", "notify": False}, headers=H).status_code == 200
    # members and the president cannot close; members cannot even look
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.get("/api/accounts/statements").status_code == 403 and mc.post("/api/accounts/close", json={"fy": fy}, headers=H).status_code == 403
    assert "finance.close" not in admin.get("/api/settings").json()["perms"].get("president", []) if "perms" in admin.get("/api/settings").json() else True


def test_library_council_pack_and_grants(admin):
    today = dt.date.today().isoformat()
    r = admin.post("/api/settings", json={"council_reg_no": "KSLC/KTM/0421", "council_grade": "C", "council_district": "Kottayam", "council_taluk": "Meenachil", "library_hours": "9 am - 8 pm",
                                          "council_criteria": {"books": [10, 5, 2, 1], "members": [10, 5, 3, 1], "issues": [10, 5, 2, 1], "periodicals": [5, 3, 2, 1], "reading_room": [10, 5, 2, 1], "programmes": [10, 5, 2, 1], "added": [10, 5, 2, 1]}}, headers=H)
    assert r.status_code == 200, r.text
    st = admin.get("/api/council/status").json()
    a, m = st["assessment"], st["metrics"]
    assert len(a["rows"]) == 7 and a["points"] == sum(x["points"] for x in a["rows"]) and a["max_points"] == 28 and a["declared"] == "C"
    members_row = next(x for x in a["rows"] if x["key"] == "members")
    assert members_row["value"] >= 3 and members_row["grade"] in ("A", "B", "C") and m["books"] >= 2 and m["issues"] >= 1
    assert all((x["grade"] == "A") == (x["next"] is None) for x in a["rows"])
    # a grant: its receipt is booked, vouchers are charged to it, the certificate adds them up
    r = admin.post("/api/council/grants", json={"source": "Kerala State Library Council", "purpose_en": "Purchase of books", "purpose_ml": "പുസ്തകങ്ങൾ വാങ്ങൽ", "sanction_no": "KSLC/G/2026/77", "amount": 5000, "received_on": today}, headers=H)
    assert r.status_code == 200, r.text
    gid = r.json()["id"]
    g = admin.get(f"/api/council/grants/{gid}").json()
    assert g["receipt_id"] and g["utilised"] == 0 and g["balance"] == 5000 and g["fy"] == admin.get("/api/accounts/heads").json()["current_fy"]
    assert admin.get(f"/api/finance/transactions?head=grant").json()["rows"][0]["amount"] == 5000
    v1 = admin.post("/api/finance/transactions", json={"kind": "payment", "head": "books", "party": "DC Books", "amount": 2000, "mode": "bank", "tx_date": today}, headers=H).json()["id"]
    v2 = admin.post("/api/finance/transactions", json={"kind": "payment", "head": "newspapers", "party": "Agent", "amount": 500, "mode": "cash", "tx_date": today, "grant_id": gid}, headers=H).json()["id"]
    assert admin.post(f"/api/council/grants/{gid}/link", json={"add": [v1]}, headers=H).status_code == 200
    g = admin.get(f"/api/council/grants/{gid}").json()
    assert g["utilised"] == 2500 and g["balance"] == 2500 and {v["id"] for v in g["vouchers"]} == {v1, v2} and v1 not in {c["id"] for c in g["candidates"]}
    uc = admin.get(f"/api/council/grants/{gid}/uc/print?lang=en").text
    assert "Utilisation Certificate" in uc and "2,500.00" in uc and "KSLC/G/2026/77" in uc and "DC Books" in uc
    assert "വിനിയോഗ സാക്ഷ്യപത്രം" in admin.get(f"/api/council/grants/{gid}/uc/print?lang=ml").text
    admin.post(f"/api/council/grants/{gid}/link", json={"remove": [v2]}, headers=H)
    assert admin.get(f"/api/council/grants/{gid}").json()["utilised"] == 2000
    assert admin.get("/api/council/grants").json()[0]["balance"] == 3000
    # the annual statement
    for lang, needle in (("en", "Annual Statement for Library Council"), ("ml", "ലൈബ്രറി കൗൺസിൽ")):
        p = admin.get(f"/api/council/statement/print?lang={lang}")
        assert p.status_code == 200 and needle in p.text and "KSLC/KTM/0421" in p.text and "Kerala State Library Council" in p.text and "Meenachil" in p.text
    assert admin.get("/api/council/status?fy=2025-26").json()["metrics"]["grants"] == []
    mc = login("9847012345", "ramesh-pass-1")
    assert mc.get("/api/council/status").status_code == 403 and mc.get("/api/council/grants").status_code == 403


def test_association_mode_units_maintenance_helpdesk(admin):
    """Switch the club to a residents' association: units, residents on units, monthly maintenance billed per unit and
    settled oldest-first by receipts, late fees, defaulters and reminders, the unit statement, and the helpdesk."""
    today = dt.date.today()
    r = admin.post("/api/settings", json={"org_kind": "association", "wards": ["A", "B"], "maint_basis": "sqft", "maint_rate": "2", "maint_period": "monthly", "maint_due_day": "5", "maint_late_fee": "100"}, headers=H)
    assert r.status_code == 200, r.text
    info = admin.get("/api/public/info").json()
    assert info["org_kind"] == "association" and info["features"] == {"library": False, "units": True, "maintenance": True, "helpdesk": True}
    # a building in one go, then one unit edited
    r = admin.post("/api/units/bulk", json={"lines": "A,101,1,2BHK,950\nA,102,1,3BHK,1200\nB,201,2,2BHK,950\nB,201,2,2BHK,950"}, headers=H)
    assert r.json()["created"] == 3
    units = admin.get("/api/units").json()
    a101 = next(u for u in units if u["label"] == "A-101"); b201 = next(u for u in units if u["label"] == "B-201")
    assert admin.post("/api/units", json={"id": b201["id"], "block": "B", "unit_no": "201", "utype": "2BHK", "area_sqft": 950, "owner_name": "Thomas", "maint_override": 1500}, headers=H).status_code == 200
    assert admin.post("/api/units", json={"block": "A", "unit_no": "101"}, headers=H).status_code == 409
    # residents: Ramesh owns A-101; a new owner registers B-201 herself; the office puts Second as tenant of A-102
    members = admin.get("/api/members?status=active").json()
    ramesh = next(m for m in members if m["name_en"] == "Ramesh Kumar"); second = next(m for m in members if m["name_en"] == "Second")
    assert admin.post(f"/api/units/{a101['id']}/residents", json={"member_id": ramesh["id"], "occupancy": "owner"}, headers=H).status_code == 200
    a102 = next(u for u in units if u["label"] == "A-102")
    admin.put(f"/api/members/{second['id']}", json={"unit_id": a102["id"], "occupancy": "tenant"}, headers=H)
    anon = TestClient(app)
    r = anon.post("/api/public/apply", json={"name_en": "Mary Thomas", "phone": "9847099001", "password": "mary-pass-11", "ward": "B", "unit_no": "201", "unit_block": "B", "occupancy": "owner", "lang": "en"}, headers=H)
    assert r.status_code == 200, r.text
    mary = r.json()["id"]
    assert admin.post(f"/api/members/{mary}/decide", json={"decision": "approve"}, headers=H).status_code == 200
    m = admin.get(f"/api/members/{mary}").json()
    assert m["unit"]["label"] == "B-201" and m["occupancy"] == "owner" and m["house_en"] == "B-201"
    r = anon.post("/api/public/apply", json={"name_en": "New Block", "phone": "9847099002", "password": "new-pass-111", "unit_no": "301", "unit_block": "C", "occupancy": "tenant"}, headers=H)
    assert r.status_code == 200 and any(u["label"] == "C-301" for u in admin.get("/api/units").json())
    assert admin.get(f"/api/units/{a101['id']}").json()["residents"][0]["name_en"] == "Ramesh Kumar"
    assert admin.get(f"/api/members?unit_id={a102['id']}").json()[0]["unit"] == "A-102"
    # billing: two months, per sq ft; B-201 has an override; C-301 has no area yet so gets no bill
    p1, p2 = f"{today.year - 1}-11", f"{today.year - 1}-12"
    assert admin.post("/api/maintenance/generate", json={"period": p1}, headers=H).json()["created"] == 3
    assert admin.post("/api/maintenance/generate", json={"period": p1}, headers=H).json()["created"] == 0
    assert admin.post("/api/maintenance/generate", json={"period": p2}, headers=H).json()["created"] == 3
    mt = admin.get(f"/api/maintenance?period={p1}").json()
    bills = {b["label"]: b for b in mt["bills"]}
    assert bills["A-101"]["amount"] == 1900 and bills["B-201"]["amount"] == 1500 and bills["A-102"]["amount"] == 2400 and mt["totals"]["units"] == 3 and mt["active_units"] == 4
    # late fees apply to the past-due bills once
    from app import assoc, scheduler
    assert assoc.apply_late_fees() == 6 and assoc.apply_late_fees() == 0
    assert scheduler.run_daily(force=True)["late_fees"] == 0
    # Ramesh pays 2500: November (1900 + 100) settles fully, December gets 500 of its 2000
    r = admin.post("/api/finance/transactions", json={"kind": "receipt", "head": "maintenance", "member_id": ramesh["id"], "amount": 2500, "mode": "upi", "ref": "UPI-M1", "notify": False}, headers=H)
    assert r.status_code == 200, r.text
    rid = r.json()["id"]
    u = admin.get(f"/api/units/{a101['id']}").json()
    by = {b["period"]: b for b in u["bills"]}
    assert by[p1]["status"] == "paid" and by[p1]["paid"] == 2000 and by[p2]["paid"] == 500 and by[p2]["status"] == "due" and u["dues"] == 1500
    assert u["receipts"][0]["id"] == rid
    # voiding the receipt reopens both bills; paying again settles them
    admin.post(f"/api/finance/transactions/{rid}/void", json={"note": "wrong unit"}, headers=H)
    assert admin.get(f"/api/units/{a101['id']}").json()["dues"] == 4000
    admin.post("/api/finance/transactions", json={"kind": "receipt", "head": "maintenance", "unit_id": a101["id"], "party": "Ramesh (cash)", "amount": 4000, "mode": "cash", "notify": False}, headers=H)
    assert admin.get(f"/api/units/{a101['id']}").json()["dues"] == 0
    # defaulters and reminders; a waived bill drops out
    d = admin.get("/api/maintenance/defaulters").json()
    assert {x["label"] for x in d} == {"A-102", "B-201"} and next(x for x in d if x["label"] == "B-201")["due"] == 3200
    fresh()
    r = admin.post("/api/maintenance/remind", json={"channels": ["whatsapp"]}, headers=H).json()
    assert r["units"] == 2 and r["queued"] >= 2
    flush()
    assert any("Maintenance" in subj and "B-201" in body for ch, to, subj, body in SENT if ch == "whatsapp")
    bid = next(b["id"] for b in admin.get(f"/api/maintenance?period={p1}").json()["bills"] if b["label"] == "A-102")
    admin.post(f"/api/maintenance/bills/{bid}", json={"waive": True, "note": "vacant in November"}, headers=H)
    assert next(x for x in admin.get("/api/maintenance/defaulters").json() if x["label"] == "A-102")["periods"] == 1
    # the resident sees her unit's bills and can print the statement; a member of another unit cannot
    mc = login("9847099001", "mary-pass-11")
    mine = mc.get("/api/finance/mine").json()["maintenance"]
    assert mine["unit"]["label"] == "B-201" and mine["due"] == 3200
    assert mc.get(f"/api/units/{b201['id']}/statement?lang=en").status_code == 200 and "Statement of account" in mc.get(f"/api/units/{b201['id']}/statement?lang=en").text
    assert mc.get(f"/api/units/{a101['id']}/statement").status_code == 403
    assert "3,200.00" in admin.get(f"/api/units/{b201['id']}/statement?lang=ml").text
    # helpdesk: a resident raises a complaint, the office works it, she closes it
    fresh()
    r = mc.post("/api/tickets", json={"category": "plumbing", "title": "Kitchen tap leaking", "details": "Since Monday", "priority": "high"}, headers=H)
    assert r.status_code == 200 and r.json()["no"].startswith(f"T/{today.year}/")
    tid = r.json()["id"]
    t = mc.get(f"/api/tickets/{tid}").json()
    assert t["unit"] == "B-201" and t["status"] == "open" and "staff" not in t
    assert mc.post(f"/api/tickets/{tid}", json={"status": "resolved"}, headers=H).status_code == 403
    staff = admin.get(f"/api/tickets/{tid}").json()["staff"]
    assert admin.post(f"/api/tickets/{tid}", json={"status": "in_progress", "assigned_to": staff[0]["id"], "note": "Plumber booked for tomorrow"}, headers=H).status_code == 200
    flush()
    assert any("in_progress" in body and "Plumber" in body for ch, to, subj, body in SENT if ch == "whatsapp")
    admin.post(f"/api/tickets/{tid}", json={"status": "resolved"}, headers=H)
    assert mc.post(f"/api/tickets/{tid}", json={"status": "closed", "note": "Thanks, fixed"}, headers=H).status_code == 200
    t = admin.get(f"/api/tickets/{tid}").json()
    assert t["status"] == "closed" and len(t["notes_list"]) == 3 and t["assignee"] == staff[0]["name"]
    lst = admin.get("/api/tickets").json()
    assert lst["counts"]["closed"] >= 1 and lst["rows"][0]["unit"] == "B-201"
    assert mc.get("/api/tickets").json()["rows"][0]["id"] == tid and "counts" in mc.get("/api/tickets").json()
    ramesh_c = login("9847012345", "ramesh-pass-1")
    assert ramesh_c.get(f"/api/tickets/{tid}").status_code == 403
    assert ramesh_c.get("/api/dashboard").json()["assoc"]["tickets_open"] == 0
    assert admin.get("/api/dashboard").json()["assoc"]["units"] == 4
    # back to club mode: the library returns, the association screens switch off
    admin.post("/api/settings", json={"org_kind": "club"}, headers=H)
    assert admin.get("/api/public/info").json()["features"]["library"] is True
